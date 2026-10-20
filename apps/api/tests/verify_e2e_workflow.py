import io
import docx
from fastapi.testclient import TestClient
from src.main import app

def run_verification():
    client = TestClient(app)

    # 1. Upload
    doc = docx.Document()
    doc.add_heading('University Assessment On Cloud Computing', level=1)
    doc.add_paragraph('Modern enterprise cloud architectures mandate multi-region failover (Smith, 2021).')
    doc.add_heading('System Methodology', level=2)
    doc.add_paragraph('We evaluate throughput using microbenchmarks.')
    doc.add_heading('References', level=1)
    doc.add_paragraph('Smith, J. (2021). Cloud Architecture Review. Tech Press.')
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    upload_res = client.post(
        '/api/jobs/upload',
        files={'file': ('assignment.docx', buf.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
    )
    assert upload_res.status_code == 200, f'Upload failed: {upload_res.text}'
    job_id = upload_res.json()['job_id']
    print('1. Uploaded project job_id:', job_id)

    # 2. Format
    format_res = client.post(f'/api/jobs/{job_id}/format', json={})
    assert format_res.status_code == 200, f'Format failed: {format_res.text}'
    print('2. Formatted successfully, current_version:', format_res.json()['current_version_id'])

    # 3. Doctor diagnose v2
    doc_res = client.get(f'/api/doctor/diagnose/{job_id}?version_id=2')
    assert doc_res.status_code == 200, f'Diagnose failed: {doc_res.text}'
    print('3. Doctor diagnosed v2, health score:', doc_res.json()['overall_health_score'])

    # 4. Upload rubric
    rubric_txt = """
Criterion 1: Architectural Soundness (40%)
Evaluate enterprise cloud architecture.
Criterion 2: Empirical Methodology (60%)
Review benchmark experiments and statistical evaluation.
"""
    rub_res = client.post(
        f'/api/jobs/{job_id}/upload-rubric',
        files={'file': ('rubric.txt', rubric_txt.encode('utf-8'), 'text/plain')},
    )
    assert rub_res.status_code == 200, f'Rubric upload failed: {rub_res.text}'
    print('4. Uploaded rubric, criteria count:', rub_res.json()['criteria_count'])

    # 5. Doctor diagnose with rubric
    doc_rub_res = client.get(f'/api/doctor/diagnose/{job_id}?version_id=2')
    assert doc_rub_res.status_code == 200
    assert doc_rub_res.json()['rubric_used'] is True
    print('5. Doctor re-diagnosed with rubric, rubric_used:', doc_rub_res.json()['rubric_used'])

    # 6. Doctor fix
    fix_res = client.post(
        f'/api/doctor/fix/{job_id}',
        json={'version_id': 2, 'selected_actions': ['repair_cover_page', 'normalize_fonts']},
    )
    assert fix_res.status_code == 200
    print('6. Doctor fixed document, new repaired file:', fix_res.json()['repaired_filename'])

    # 7. Plan project slides
    plan_res = client.post(f'/api/presentation/project/{job_id}/plan', json={'target_slide_count': 6})
    assert plan_res.status_code == 200
    outline = plan_res.json()
    print('7. Planned slides from v3, slide count:', len(outline['slides']))

    # 8. Generate project presentation
    gen_res = client.post(
        f'/api/presentation/project/{job_id}/generate',
        json={'outline': outline, 'theme': 'premium_university'},
    )
    assert gen_res.status_code == 200
    print('8. Generated presentation:', gen_res.json()['filename'], 'is_stale:', gen_res.json()['is_stale'])

    # 9. Rubric analysis endpoint
    rub_ana_res = client.get(f'/api/jobs/{job_id}/rubric-analysis')
    assert rub_ana_res.status_code == 200
    print('9. Rubric analysis retrieved without re-upload, compliance_score:', rub_ana_res.json()['overall_compliance_score'])

    # 10. Check version list
    v_res = client.get(f'/api/jobs/{job_id}/versions')
    assert v_res.status_code == 200
    versions = v_res.json()['versions']
    assert len(versions) >= 3
    print('10. Document versions tracked:', [f"v{v['version_id']}: {v['version_label']}" for v in versions])

    print('\n>>> ALL 10 END-TO-END STEPS VERIFIED 100% PASS! <<<')

if __name__ == '__main__':
    run_verification()
