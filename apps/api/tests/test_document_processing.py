import pytest
from pathlib import Path
import docx
from docx.shared import Inches, Pt
from src.modules.jobs.job_schema import (
    FormattingSettings,
    MarginSettings,
    BodyStyle,
    HeadingStyle,
    CoverPageData,
    MarkingSchemeData,
    MarkingCriterion,
)
from src.modules.documents.formatter import apply_document_formatting
from src.modules.documents.cover_page import apply_cover_page
from src.modules.documents.marking_scheme import apply_marking_scheme
from src.modules.documents.page_numbering import apply_page_numbering
from src.modules.documents.compliance import check_compliance
from src.modules.documents.inspector import inspect_docx
from src.common.security import validate_docx_bytes, get_job_directory, validate_uuid
from src.common.errors import FileValidationError

@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    doc_path = tmp_path / "sample_assignment.docx"
    doc = docx.Document()

    # Add Title & Heading 1
    h1 = doc.add_paragraph("Introduction", style="Heading 1")
    p1 = doc.add_paragraph("This is the introductory paragraph analyzing modern distributed architectures.")
    run_bold = p1.add_run(" Bold keyword here.")
    run_bold.bold = True

    # Heading 2
    h2 = doc.add_paragraph("Background Literature", style="Heading 2")
    p2 = doc.add_paragraph("Extensive literature review discussing consensus protocols and cloud systems.")

    # Table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Latency"
    table.rows[1].cells[1].text = "12ms"

    # Conclusion & References
    doc.add_paragraph("Conclusion", style="Heading 1")
    doc.add_paragraph("In conclusion, the system meets the high-availability performance criteria.")

    doc.add_paragraph("References", style="Heading 1")
    doc.add_paragraph("1. Lamport, L. (1998). The Part-Time Parliament. ACM TOCS, 16(2), 133-169.")

    doc.save(doc_path)
    return doc_path

def test_inspect_docx(sample_docx: Path):
    res = inspect_docx(sample_docx)
    assert res.total_paragraphs >= 5
    assert res.total_tables == 1
    assert any(h.text == "Introduction" for h in res.detected_headings)
    assert res.has_references_section is True

def test_apply_formatting(sample_docx: Path):
    settings = FormattingSettings(
        paper_size="A4",
        margins=MarginSettings(top_inches=1.0, bottom_inches=1.0, left_inches=1.0, right_inches=1.0),
        body=BodyStyle(font_family="Times New Roman", font_size_pt=12.0, line_spacing=1.5),
        heading1=HeadingStyle(font_family="Times New Roman", font_size_pt=16.0, bold=True, page_break_before=True),
        start_main_heading_on_new_page=True,
    )

    doc = docx.Document(sample_docx)
    apply_document_formatting(doc, settings)
    doc.save(sample_docx)

    # Reopen and verify
    reloaded = docx.Document(sample_docx)
    sec = reloaded.sections[0]
    assert round(sec.top_margin.inches, 1) == 1.0
    assert round(sec.left_margin.inches, 1) == 1.0

def test_cover_page_and_marking_scheme(sample_docx: Path):
    settings = FormattingSettings(
        cover_page=CoverPageData(
            enabled=True,
            student_name="Jane Doe",
            student_id="U12345678",
            institution="Oxford University",
            department="Department of Computer Science",
            assignment_title="Distributed Systems Analysis",
            course="CS401 Distributed Computing",
            lecturer_name="Dr. Alan Turing",
            submission_date="2026-10-20",
        ),
        marking_scheme=MarkingSchemeData(
            enabled=True,
            title="Module Marking Scheme",
            criteria=[
                MarkingCriterion(name="Core Implementation", description="Soundness of algorithms", max_marks=50.0),
                MarkingCriterion(name="Critical Evaluation", description="Empirical benchmark analysis", max_marks=50.0),
            ]
        ),
    )

    doc = docx.Document(sample_docx)
    apply_cover_page(doc, settings.cover_page, settings)
    apply_marking_scheme(doc, settings.marking_scheme, settings)
    apply_page_numbering(doc, settings)
    doc.save(sample_docx)

    # Run compliance
    compliance = check_compliance(sample_docx, settings)
    assert compliance.score_percentage > 70
    assert any(c.id in ["cover-page", "cover-page-check"] and c.status == "PASS" for c in compliance.checks)

def test_security_validation():
    # Invalid magic bytes should raise FileValidationError
    fake_content = b"This is not a zip or docx file"
    with pytest.raises(FileValidationError):
        validate_docx_bytes(fake_content, "test.docx")

    # Invalid UUID
    with pytest.raises(FileValidationError):
        validate_uuid("invalid-uuid-123")
