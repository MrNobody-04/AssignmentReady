import io
import pptx
from src.modules.presentation.ppt_generator import (
    PresentationOutline,
    SlideItem,
    MetricItem,
    THEMES,
    generate_powerpoint_presentation,
    classify_slide_type,
)

def test_all_eight_presentation_themes():
    sample_outline = PresentationOutline(
        title="Comparative Distributed Systems",
        subtitle="Performance, Consensus & Fault Tolerance",
        author="Academic Researcher",
        institution="Faculty of Engineering",
        slides=[
            SlideItem(title="Presentation Agenda", slide_type="agenda", bullets=["Introduction", "Consensus Models", "Evaluation"]),
            SlideItem(title="Project Objectives", slide_type="objectives", bullets=["Low latency", "High throughput", "Partition tolerance"]),
            SlideItem(title="Execution Workflow", slide_type="process_workflow", bullets=["Ingest telemetry", "Run Raft consensus", "Commit log state"]),
            SlideItem(title="Paxos vs Raft", slide_type="comparison", column_left_title="Multi-Paxos", column_right_title="Raft Algorithm"),
            SlideItem(title="Empirical Results", slide_type="metrics_stats", metrics=[
                MetricItem(value="99.99%", label="Uptime Reliability"),
                MetricItem(value="4.5ms", label="Average Latency"),
            ]),
            SlideItem(title="Benchmark Matrix", slide_type="table", table_headers=["Metric", "Result"], table_rows=[["Latency", "4.5ms"]]),
            SlideItem(title="Summary Conclusions", slide_type="conclusion", bullets=["Raft maintains formal correctness while simplifying leader election."]),
            SlideItem(title="Academic References", slide_type="references", bullets=["[1] Ongaro, D., & Ousterhout, J. (2014). In Search of an Understandable Consensus Algorithm."]),
            SlideItem(title="Questions & Answers", slide_type="qa_closing"),
        ],
    )

    theme_keys = [
        "modern_academic",
        "tech_innovation",
        "minimalist_premium",
        "corporate_pro",
        "creative_gradient",
        "research_science",
        "dark_presentation",
        "university_classic",
    ]

    for key in theme_keys:
        assert key in THEMES, f"Theme '{key}' must exist in THEMES"
        pptx_bytes = generate_powerpoint_presentation(sample_outline, theme_key=key)
        assert len(pptx_bytes) > 10000, f"Theme '{key}' generated empty or corrupted deck"

        # Reopen with python-pptx to verify integrity
        prs = pptx.Presentation(io.BytesIO(pptx_bytes))
        assert len(prs.slides) == len(sample_outline.slides) + 1, f"Expected {len(sample_outline.slides) + 1} slides for theme {key}"
        assert abs(prs.slide_width.inches - 13.333) < 0.01, "Expected 16:9 widescreen width"
        assert abs(prs.slide_height.inches - 7.5) < 0.01, "Expected 16:9 widescreen height"

def test_speaker_notes_preservation():
    outline = PresentationOutline(
        title="Speaker Notes Validation",
        slides=[
            SlideItem(
                title="Methodology Overview",
                slide_type="content",
                bullets=["Sample size: 5,000 requests", "Measured over 72 hours"],
                notes="Remember to emphasize that the benchmark was conducted on isolated physical nodes without virtualization overhead.",
            )
        ],
    )

    pptx_bytes = generate_powerpoint_presentation(outline, theme_key="modern_academic")
    prs = pptx.Presentation(io.BytesIO(pptx_bytes))
    content_slide = prs.slides[1]
    assert content_slide.has_notes_slide
    notes_text = content_slide.notes_slide.notes_text_frame.text
    assert "isolated physical nodes" in notes_text

def test_intelligent_slide_classification():
    assert classify_slide_type("Presentation Agenda & Topics", []) == "agenda"
    assert classify_slide_type("Core Objectives & Aims", []) == "objectives"
    assert classify_slide_type("Methodology & Experimental Procedure", []) == "process_workflow"
    assert classify_slide_type("Architectural Comparison & Tradeoffs", []) == "comparison"
    assert classify_slide_type("Quantitative Performance Metrics", []) == "metrics_stats"
    assert classify_slide_type("Summary & Key Takeaways", []) == "conclusion"
    assert classify_slide_type("Academic Bibliography & References", []) == "references"
    assert classify_slide_type("Questions & Open Discussion", []) == "qa_closing"
