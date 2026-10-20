import io
from pathlib import Path
import docx
from fastapi.testclient import TestClient
from src.main import app
from src.modules.references.citation_analyzer import analyze_document_citations
from src.modules.doctor.doctor_engine import diagnose_document, DoctorIssue
from src.modules.jobs.job_schema import FormattingSettings

client = TestClient(app)

def create_sample_academic_docx(tmp_path: Path) -> Path:
    doc = docx.Document()
    # Add title and unstandardized body font
    doc.add_heading("Deep Learning in Autonomous Navigation", level=1)
    p1 = doc.add_paragraph("Machine learning systems have gained widespread adoption (Smith, 2021). Autonomous agents rely on deep vision [1].")
    p2 = doc.add_paragraph("Recent empirical benchmarks indicate substantial latency reductions (Johnson & Davis, 2022).")
    
    # Subheading jump: skip H2 and go straight to H3
    doc.add_heading("Sub-System Benchmarks", level=3)
    p3 = doc.add_paragraph("Data throughput metrics were recorded across multiple edge GPU devices (UnknownAuthor, 2024).")

    # References section
    doc.add_heading("References", level=1)
    doc.add_paragraph("Smith, J. (2021). Foundations of Vision-Based Navigation Systems. Academic Robotics Press.")
    doc.add_paragraph("Johnson, R., & Davis, M. (2022). Empirical Benchmarks for Autonomous Edge Devices. Journal of AI.")
    doc.add_paragraph("UncitedPaper, A. (2019). Unused Methodology. Journal of Unused Studies.")

    path = tmp_path / "test_sample_paper.docx"
    doc.save(str(path))
    return path

def test_citation_analyzer(tmp_path):
    path = create_sample_academic_docx(tmp_path)
    report = analyze_document_citations(path)

    assert report.total_in_text_citations >= 3
    assert report.total_references_entries == 3
    # Both author-date and numeric [1] were inserted, so style should be mixed
    assert report.detected_style == "Mixed / Inconsistent"
    assert len(report.style_inconsistencies) > 0
    # UnknownAuthor, 2024 is in body but not in bibliography
    unmatched_names = [c.author for c in report.unmatched_in_text_citations]
    assert any("UnknownAuthor" in name for name in unmatched_names)
    # UncitedPaper is in bibliography but never cited in text
    uncited_texts = [r.raw_text for r in report.uncited_reference_entries]
    assert any("UncitedPaper" in t for t in uncited_texts)

def test_doctor_diagnose_and_fix_api(tmp_path):
    path = create_sample_academic_docx(tmp_path)
    with open(path, "rb") as f:
        file_bytes = f.read()

    # 1. Test diagnose-upload endpoint
    res = client.post(
        "/api/doctor/diagnose-upload",
        files={"file": ("paper.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    diag = res.json()
    assert "job_id" in diag
    job_id = diag["job_id"]
    assert diag["total_issues"] > 0
    # Should detect heading jump, mixed citations, unmatched citation
    issue_ids = [i["id"] for i in diag["issues"]]
    assert "doc-heading-jump" in issue_ids or "doc-citation-style-inconsistent" in issue_ids

    # 2. Test fix endpoint
    fix_res = client.post(
        f"/api/doctor/fix/{job_id}",
        json={"fix_all_safe": True, "selected_action_ids": ["standardize_body_font", "apply_standard_margins", "inject_cover_page"]},
    )
    assert fix_res.status_code == 200
    fix_data = fix_res.json()
    assert fix_data["job_id"] == job_id
    assert "download_url" in fix_data
    assert "repaired_filename" in fix_data
