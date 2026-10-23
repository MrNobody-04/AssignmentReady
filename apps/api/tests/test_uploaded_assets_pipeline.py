from pathlib import Path

import docx

from src.modules.documents.formatter import apply_formatting
from src.modules.documents.guideline_analyzer import analyze_guidelines
from src.modules.jobs.job_schema import FormattingSettings, MarkingCriterion
from src.modules.presentation.ppt_generator import PresentationOutline
from src.modules.presentation.presentation_controller import _apply_project_cover_metadata


def test_guideline_analysis_returns_reviewable_formatting_suggestions():
    analysis = analyze_guidelines(
        "Use A4 paper, Times New Roman font size 12 pt, 1.5 line spacing, "
        "1 inch margins. Include Introduction, Methodology, Conclusion and References."
    )

    assert analysis["has_suggestions"] is True
    assert analysis["suggestions"]["paper_size"] == "A4"
    assert analysis["suggestions"]["body"]["font_family"] == "Times New Roman"
    assert analysis["suggestions"]["body"]["line_spacing"] == 1.5
    assert "References" in analysis["suggestions"]["required_sections"]


def test_uploaded_template_and_after_cover_rubric_are_preserved(tmp_path: Path):
    source = docx.Document()
    source.add_heading("Introduction", level=1)
    source.add_paragraph("Assignment body content.")
    source.add_heading("References", level=1)
    source_path = tmp_path / "source.docx"
    source.save(source_path)

    template = docx.Document()
    template.add_paragraph("OFFICIAL FACULTY COVER")
    template.add_paragraph("Candidate: {{student_name}}")
    template_table = template.add_table(rows=1, cols=2)
    template_table.cell(0, 0).text = "Module"
    template_table.cell(0, 1).text = "{{module_code}}"
    template_path = tmp_path / "cover.docx"
    template.save(template_path)

    settings = FormattingSettings()
    settings.cover_page.mode = "template"
    settings.cover_page.fill_placeholders = True
    settings.cover_page.student_name = "Asha Rai"
    settings.cover_page.module_code = "CS401"
    settings.marking_scheme.enabled = True
    settings.marking_scheme.placement = "after_cover"
    settings.marking_scheme.criteria = [
        MarkingCriterion(name="Analysis", description="Critical evaluation", max_marks=50, weight_percentage=50)
    ]
    output = tmp_path / "formatted.docx"

    apply_formatting(source_path, output, settings, cover_template_path=template_path)

    result = docx.Document(output)
    text = "\n".join(paragraph.text for paragraph in result.paragraphs)
    table_text = "\n".join(cell.text for table in result.tables for row in table.rows for cell in row.cells)
    assert "OFFICIAL FACULTY COVER" in text
    assert "Candidate: Asha Rai" in text
    assert "CS401" in table_text
    assert "Analysis" in table_text


def test_project_presentation_uses_approved_cover_metadata():
    settings = FormattingSettings()
    settings.cover_page.assignment_title = "Distributed Systems Evaluation"
    settings.cover_page.student_name = "Asha Rai"
    settings.cover_page.student_id = "CS-104"
    settings.cover_page.institution = "Example University"
    settings.cover_page.department = "School of Computing"
    settings.cover_page.course = "BSc Computer Science"
    settings.cover_page.module_code = "CS401"
    settings.cover_page.lecturer_name = "Dr. Shah"
    settings.cover_page.submission_date = "20 October 2026"

    outline = _apply_project_cover_metadata(PresentationOutline(), settings)

    assert outline.title == "Distributed Systems Evaluation"
    assert outline.author == "Asha Rai"
    assert outline.college_name == "School of Computing"
    assert outline.course_code == "CS401"
    assert outline.lecturer_name == "Dr. Shah"
