import tempfile
from pathlib import Path
import docx
from docx.oxml.ns import qn
from src.modules.jobs.job_schema import (
    FormattingSettings,
    CoverPageData,
    DeclarationSettings,
    AcknowledgementSettings,
    AbstractSettings,
    TableOfContentsSettings,
    ListOfFiguresSettings,
    ListOfTablesSettings,
    ListOfAbbreviationsSettings,
    AbbreviationItem,
    PageNumberingSettings,
)
from src.modules.documents.formatter import apply_formatting
from src.modules.documents.inspector import inspect_docx

def create_sample_academic_doc(file_path: Path) -> None:
    doc = docx.Document()
    # Heading 1
    doc.add_heading("Chapter 1: Introduction to Cloud Security", level=1)
    doc.add_paragraph("Cloud infrastructure security involves deterministic access controls and data isolation.")

    # Figure 1 with caption
    doc.add_paragraph("[Image: Architecture]")
    doc.add_paragraph("Figure 1: Enterprise Microservices Architecture")

    # Heading 2
    doc.add_heading("1.1 Threat Modeling", level=2)
    doc.add_paragraph("Threat models must evaluate both insider risks and external lateral movement.")

    # Table 1 with caption
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Risk"
    table.rows[0].cells[1].text = "Severity"
    table.rows[1].cells[0].text = "Data Leak"
    table.rows[1].cells[1].text = "High"
    doc.add_paragraph("Table 1: Identified Security Vulnerabilities")

    # Figure 2 with Fig. abbreviation
    doc.add_paragraph("Fig. 2: Network Latency Distribution")

    # Conclusion & References
    doc.add_heading("Chapter 2: Conclusion", level=1)
    doc.add_paragraph("The proposed zero-trust framework successfully mitigates unauthorized privilege escalation.")

    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Smith, J. (2025). Cloud Security Fundamentals. Academic Press.")

    doc.save(str(file_path))

def test_full_academic_front_matter_and_sections():
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        input_file = tmp_path / "sample_input.docx"
        output_file = tmp_path / "formatted_output.docx"

        create_sample_academic_doc(input_file)

        # Pre-inspection check
        inspect_res = inspect_docx(input_file)
        assert len(inspect_res.detected_figures) == 2
        assert inspect_res.detected_figures[0].number == "1"
        assert "Architecture" in inspect_res.detected_figures[0].caption
        assert inspect_res.detected_figures[1].number == "2"

        assert len(inspect_res.detected_table_captions) == 1
        assert inspect_res.detected_table_captions[0].number == "1"

        settings = FormattingSettings(
            cover_page=CoverPageData(
                enabled=True,
                student_name="Test Student",
                student_id="ST123456",
                assignment_title="Advanced Cloud Security",
                institution="University of Technology",
            ),
            declaration=DeclarationSettings(enabled=True),
            acknowledgement=AcknowledgementSettings(enabled=True),
            abstract=AbstractSettings(enabled=True),
            table_of_contents=TableOfContentsSettings(enabled=True),
            list_of_figures=ListOfFiguresSettings(enabled=True, only_if_figures_exist=True),
            list_of_tables=ListOfTablesSettings(enabled=True, only_if_tables_exist=True),
            list_of_abbreviations=ListOfAbbreviationsSettings(
                enabled=True,
                items=[
                    AbbreviationItem(term="CSP", definition="Cloud Service Provider"),
                    AbbreviationItem(term="IAM", definition="Identity and Access Management"),
                ],
            ),
            page_numbering=PageNumberingSettings(
                enabled=True,
                hide_on_cover=True,
                roman_front_matter=True,
                restart_at_introduction=True,
            ),
        )

        apply_formatting(input_file, output_file, settings)
        assert output_file.exists()

        # Verify output structure and sections
        res_doc = docx.Document(output_file)
        sections = res_doc.sections
        assert len(sections) == 3, f"Expected 3 sections (Cover, Front Matter, Body), got {len(sections)}"

        # Section 0: Cover
        sec0 = sections[0]
        assert sec0.footer.is_linked_to_previous is False
        for p in sec0.footer.paragraphs:
            assert p.text.strip() == "", "Cover footer should have no visible text"

        # Section 1: Front Matter (Roman numerals i, ii, iii)
        sec1 = sections[1]
        assert sec1.footer.is_linked_to_previous is False
        pg1 = sec1._sectPr.find(qn("w:pgNumType"))
        assert pg1 is not None
        assert pg1.get(qn("w:fmt")) == "lowerRoman"
        assert pg1.get(qn("w:start")) == "1"

        # Section 2: Main Body (Arabic numerals 1, 2, 3)
        sec2 = sections[2]
        assert sec2.footer.is_linked_to_previous is False
        pg2 = sec2._sectPr.find(qn("w:pgNumType"))
        assert pg2 is not None
        assert pg2.get(qn("w:fmt")) == "decimal"
        assert pg2.get(qn("w:start")) == "1"

        # Verify order of text headings in document
        all_headings = [p.text.strip() for p in res_doc.paragraphs if p.style and "Heading 1" in p.style.name]
        expected_front_matter = [
            "Declaration of Originality",
            "Acknowledgements",
            "Abstract",
            "Table of Contents",
            "List of Figures",
            "List of Tables",
            "List of Abbreviations",
            "Chapter 1: Introduction to Cloud Security",
            "Chapter 2: Conclusion",
            "References",
        ]
        assert all_headings == expected_front_matter

def test_suppress_empty_lof_and_lot():
    """Verify LOF and LOT are omitted when document has no figures and no tables."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        input_file = tmp_path / "plain_doc.docx"
        output_file = tmp_path / "plain_output.docx"

        doc = docx.Document()
        doc.add_heading("Chapter 1: Overview", level=1)
        doc.add_paragraph("Plain text with no figures or tables.")
        doc.save(str(input_file))

        settings = FormattingSettings(
            cover_page=CoverPageData(enabled=False),
            table_of_contents=TableOfContentsSettings(enabled=True),
            list_of_figures=ListOfFiguresSettings(enabled=True, only_if_figures_exist=True),
            list_of_tables=ListOfTablesSettings(enabled=True, only_if_tables_exist=True),
        )

        apply_formatting(input_file, output_file, settings)
        res_doc = docx.Document(output_file)

        headings = [p.text.strip() for p in res_doc.paragraphs if p.style and "Heading 1" in p.style.name]
        assert "Table of Contents" in headings
        assert "List of Figures" not in headings, "LOF should be suppressed when no figures exist"
        assert "List of Tables" not in headings, "LOT should be suppressed when no tables exist"
