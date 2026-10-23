import pytest
from pathlib import Path
import docx

from src.modules.jobs.job_schema import CoverPageData
from src.modules.documents.composer import (
    compose_submission_document,
    replace_cover_placeholders,
)


def _create_template_with_placeholders(path: Path):
    doc = docx.Document()
    p1 = doc.add_paragraph()
    r1 = p1.add_run("UNIVERSITI TEKNOLOGI - FACULTY OF COMPUTING")
    r1.bold = True

    p2 = doc.add_paragraph()
    r2_pre = p2.add_run("Candidate Name: ")
    r2_target = p2.add_run("[Student Name]")
    r2_target.italic = True

    p3 = doc.add_paragraph()
    p3.add_run("Student ID: ")
    p3.add_run("{{student_id}}")

    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Module Title:"
    table.cell(0, 1).text = "[Assignment Title]"

    doc.save(str(path))


def _create_dummy_body(path: Path):
    doc = docx.Document()
    doc.add_heading("1.0 Introduction", level=1)
    doc.add_paragraph("Body paragraph text.")
    doc.save(str(path))


def test_placeholder_replacement_off_by_default(tmp_path: Path):
    template_path = tmp_path / "cover_template.docx"
    _create_template_with_placeholders(template_path)

    body_path = tmp_path / "body.docx"
    _create_dummy_body(body_path)

    output_path = tmp_path / "output_default.docx"

    # Default settings: fill_placeholders = False
    cover_data = CoverPageData(
        student_name="Alexander Hamilton",
        student_id="AH-1789",
        assignment_title="Treasury Financial Systems",
    )
    assert cover_data.fill_placeholders is False
    assert cover_data.use_exact_cover is True

    compose_submission_document(
        body_docx_path=body_path,
        output_docx_path=output_path,
        cover_path=template_path,
        cover_page_data=cover_data,
    )

    out_doc = docx.Document(str(output_path))
    full_text = "\n".join(p.text for p in out_doc.paragraphs)
    table_text = "\n".join(c.text for t in out_doc.tables for row in t.rows for c in row.cells)

    # UNCHANGED rule: Placeholders must NOT be replaced when fill_placeholders is False
    assert "[Student Name]" in full_text
    assert "{{student_id}}" in full_text
    assert "[Assignment Title]" in table_text
    assert "Alexander Hamilton" not in full_text


def test_placeholder_replacement_when_explicitly_enabled(tmp_path: Path):
    template_path = tmp_path / "cover_template.docx"
    _create_template_with_placeholders(template_path)

    body_path = tmp_path / "body.docx"
    _create_dummy_body(body_path)

    output_path = tmp_path / "output_enabled.docx"

    # User explicitly enables "Fill in cover details"
    cover_data = CoverPageData(
        fill_placeholders=True,
        student_name="Alexander Hamilton",
        student_id="AH-1789",
        assignment_title="Treasury Financial Systems",
    )
    assert cover_data.fill_placeholders is True

    compose_submission_document(
        body_docx_path=body_path,
        output_docx_path=output_path,
        cover_path=template_path,
        cover_page_data=cover_data,
    )

    out_doc = docx.Document(str(output_path))
    full_text = "\n".join(p.text for p in out_doc.paragraphs)
    table_text = "\n".join(c.text for t in out_doc.tables for row in t.rows for c in row.cells)

    # When enabled, placeholders ARE replaced
    assert "Alexander Hamilton" in full_text
    assert "AH-1789" in full_text
    assert "Treasury Financial Systems" in table_text

    # And raw placeholder tokens should be gone
    assert "[Student Name]" not in full_text
    assert "{{student_id}}" not in full_text
    assert "[Assignment Title]" not in table_text
