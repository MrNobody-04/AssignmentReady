import os
import re
import pytest
from pathlib import Path
import docx
import pypdf

from src.modules.jobs.job_schema import FormattingSettings, TableOfContentsSettings
from src.modules.documents.formatter import apply_formatting
from src.modules.jobs.job_manager import job_manager
from src.modules.documents.word_updater import is_word_com_available


def _create_multipage_assignment(path: Path):
    doc = docx.Document()

    # Introduction
    h1 = doc.add_heading("1.0 Introduction", level=1)
    for _ in range(4):
        doc.add_paragraph(
            "This introductory section provides context for the enterprise resource planning implementation. "
            "It outlines key stakeholder goals, operational efficiency requirements, and system integration strategies."
        )

    # Force page break to guarantee Section 2 is on next page
    doc.add_page_break()

    # Methodology
    doc.add_heading("2.0 Enterprise Architecture Methodology", level=1)
    for _ in range(5):
        doc.add_paragraph(
            "The methodology incorporates cross-functional process alignment, database normalization, "
            "and continuous data pipeline orchestration across distributed cloud systems."
        )

    # Force page break to guarantee Section 3 is on next page
    doc.add_page_break()

    # Results & Conclusion
    doc.add_heading("3.0 Empirical Findings and Strategic Outcomes", level=1)
    for _ in range(4):
        doc.add_paragraph(
            "The empirical results demonstrate a 42% reduction in transaction latency, automated reconciliation, "
            "and zero downtime across all departmental units."
        )

    doc.save(str(path))


def test_update_fields_xml_flag_in_docx(tmp_path: Path):
    input_doc = tmp_path / "input.docx"
    _create_multipage_assignment(input_doc)

    output_doc = tmp_path / "formatted.docx"
    settings = FormattingSettings()
    settings.table_of_contents.enabled = True
    settings.heading1.page_break_before = True

    apply_formatting(input_doc, output_doc, settings)
    assert output_doc.exists()

    # Inspect settings.xml inside DOCX zip
    import zipfile
    with zipfile.ZipFile(str(output_doc), "r") as z:
        settings_xml = z.read("word/settings.xml").decode("utf-8")
        assert "w:updateFields" in settings_xml
        assert 'w:val="true"' in settings_xml or "w:val='true'" in settings_xml


def test_toc_page_numbers_in_exported_pdf(tmp_path: Path):
    input_doc = tmp_path / "input.docx"
    _create_multipage_assignment(input_doc)

    output_doc = tmp_path / "formatted.docx"
    settings = FormattingSettings()
    settings.table_of_contents.enabled = True
    settings.heading1.page_break_before = True

    apply_formatting(input_doc, output_doc, settings)
    assert output_doc.exists()

    # Convert to PDF
    pdf_out = job_manager._convert_to_pdf(output_doc, tmp_path)
    if not pdf_out or not pdf_out.exists():
        pytest.skip("PDF conversion engine (Word COM or LibreOffice) not active in this test environment.")

    reader = pypdf.PdfReader(str(pdf_out))
    num_pages = len(reader.pages)
    assert num_pages >= 3, f"Expected multi-page PDF, got {num_pages} pages"

    # Check that TOC (Page 2) contains entries
    toc_page_text = reader.pages[1].extract_text()
    assert "Table of Contents" in toc_page_text
    assert "1.0 Introduction" in toc_page_text
    assert "2.0 Enterprise Architecture" in toc_page_text

    # Locate which physical PDF page contains the body headings (pages after TOC)
    heading_page_map = {}
    for p_idx in range(2, num_pages):
        text = reader.pages[p_idx].extract_text()
        if "1.0 Introduction" in text:
            heading_page_map.setdefault("1.0 Introduction", p_idx + 1)
        if "2.0 Enterprise Architecture Methodology" in text:
            heading_page_map.setdefault("2.0 Enterprise Architecture", p_idx + 1)
        if "3.0 Empirical Findings" in text:
            heading_page_map.setdefault("3.0 Empirical Findings", p_idx + 1)

    # All headings should be placed on sequential physical body pages
    assert "1.0 Introduction" in heading_page_map
    assert "2.0 Enterprise Architecture" in heading_page_map
    assert "3.0 Empirical Findings" in heading_page_map
    assert heading_page_map["2.0 Enterprise Architecture"] > heading_page_map["1.0 Introduction"]
    assert heading_page_map["3.0 Empirical Findings"] > heading_page_map["2.0 Enterprise Architecture"]
