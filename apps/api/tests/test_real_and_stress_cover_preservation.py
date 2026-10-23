import os
import io
import shutil
import hashlib
from pathlib import Path
import pytest
import docx
import pypdf

from src.modules.jobs.job_schema import CoverPageData
from src.modules.documents.composer import (
    compose_submission_document,
    compose_submission_pdf,
    compare_docx_cover_parts,
    compare_rendered_cover_pixels,
)

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures" / "real_covers"


def _create_sample_body(dest_path: Path):
    doc = docx.Document()
    doc.add_heading("1.0 Introduction", level=1)
    doc.add_paragraph("This is the introductory paragraph of the academic submission.")
    doc.add_heading("2.0 Architecture & Methodology", level=1)
    doc.add_paragraph("Enterprise Resource Planning systems integrate core functional areas.")
    doc.add_heading("3.0 Conclusion", level=1)
    doc.add_paragraph("In summary, the ERP implementation yielded significant business value.")
    doc.save(str(dest_path))


def test_fixtures_directory_exists_and_has_covers():
    assert FIXTURES_DIR.exists(), f"Fixtures directory missing: {FIXTURES_DIR}"
    files = list(FIXTURES_DIR.glob("*.*"))
    assert len(files) >= 5, f"Expected at least 5 test covers, found {len(files)}"


@pytest.mark.parametrize("cover_filename", [
    "nilai_university_cover.docx",
    "stress_textbox_shapes.docx",
    "stress_content_controls_sdt.docx",
    "stress_nested_tables.docx",
    "stress_landscape_sections.docx",
    "stress_header_footer_media.docx",
])
def test_real_and_stress_docx_cover_preservation(cover_filename: str, tmp_path: Path):
    cover_path = FIXTURES_DIR / cover_filename
    if not cover_path.exists():
        pytest.skip(f"Fixture {cover_filename} not found in {FIXTURES_DIR}")

    body_path = tmp_path / "body.docx"
    _create_sample_body(body_path)

    output_path = tmp_path / f"output_{cover_filename}"
    cover_data = CoverPageData(
        enabled=True,
        mode="template",
        use_exact_cover=True,
        fill_placeholders=False,
    )

    result = compose_submission_document(
        body_docx_path=body_path,
        output_docx_path=output_path,
        cover_path=cover_path,
        cover_page_data=cover_data,
    )

    assert result["status"] == "success"
    assert output_path.exists()

    # Part-by-part OOXML comparison
    comp = compare_docx_cover_parts(cover_path, output_path)
    if not comp["identical"]:
        pytest.fail(f"Cover preservation failed for {cover_filename}! Differing parts: {comp['differing_parts']}, details: {comp['details']}")

    assert comp["identical"] is True


@pytest.mark.parametrize("cover_filename", [
    "nilai_university_cover.pdf",
    "official_vector_cover.pdf",
])
def test_real_and_stress_pdf_cover_preservation(cover_filename: str, tmp_path: Path):
    cover_path = FIXTURES_DIR / cover_filename
    if not cover_path.exists():
        pytest.skip(f"Fixture {cover_filename} not found in {FIXTURES_DIR}")

    # Create dummy body PDF
    body_pdf = tmp_path / "body.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with open(body_pdf, "wb") as f:
        writer.write(f)

    output_pdf = tmp_path / f"output_{cover_filename}"
    compose_submission_pdf(
        cover_pdf_path=cover_path,
        body_pdf_path=body_pdf,
        output_pdf_path=output_pdf,
    )

    assert output_pdf.exists()

    # Verify Page 1 of output is byte-identical or structurally identical to original cover
    orig_reader = pypdf.PdfReader(str(cover_path))
    out_reader = pypdf.PdfReader(str(output_pdf))

    assert len(out_reader.pages) >= 2
    orig_p1_text = orig_reader.pages[0].extract_text()
    out_p1_text = out_reader.pages[0].extract_text()
    assert orig_p1_text == out_p1_text

    # Visual pixel diff check
    passed, ratio, details = compare_rendered_cover_pixels(cover_path, output_pdf, max_diff_ratio=0.005)
    assert passed, f"Pixel difference on {cover_filename} exceeded threshold: {details}"
