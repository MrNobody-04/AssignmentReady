import io
import docx
import pytest
from fastapi.testclient import TestClient

from src.main import app
from src.modules.jobs.job_schema import FormattingSettings

client = TestClient(app)

def create_sample_docx(title="Distributed Systems Research Paper", author="Jane Doe"):
    doc = docx.Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Author: {author}")
    
    doc.add_heading("1. Introduction & Background", level=1)
    doc.add_paragraph("Distributed systems require resilient consensus protocols and decoupled messaging (Lamport, 1978).")
    doc.add_paragraph("Modern architectures deploy Byzantine fault tolerance across untrusted compute nodes.")
    
    doc.add_heading("2. System Architecture & Methodology", level=1)
    doc.add_paragraph("Our empirical evaluation utilizes microbenchmarks across 16 geo-distributed clusters.")
    doc.add_paragraph("Table 1 summarizes cluster specifications.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Node ID"
    table.cell(0, 1).text = "VCPU / RAM"
    table.cell(1, 0).text = "node-alpha"
    table.cell(1, 1).text = "8 vCPU / 32GB"
    
    doc.add_heading("3. Performance Analysis & Results", level=1)
    doc.add_paragraph("Latency was reduced by 34.2% while throughput reached 125,000 transactions per second.")
    
    doc.add_heading("4. Conclusion", level=1)
    doc.add_paragraph("The consensus protocol guarantees safety and deterministic liveness under partial network partitions.")
    
    doc.add_heading("References", level=1)
    doc.add_paragraph("Lamport, L. (1978). Time, clocks, and the ordering of events in a distributed system. CACM, 21(7), 558-565.")
    
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def create_sample_rubric_docx():
    doc = docx.Document()
    doc.add_heading("Assessment Brief & Marking Scheme (CS401)", level=1)
    doc.add_paragraph("Criterion 1: Introduction & Background Analysis (20 Marks)\nClear context, motivation, and theoretical grounding.")
    doc.add_paragraph("Criterion 2: System Architecture & Methodology (30 Marks)\nRigorous architecture design and experimental methodology.")
    doc.add_paragraph("Criterion 3: Performance Analysis & Evaluation (40 Marks)\nEmpirical performance results, benchmarking, and throughput graphs.")
    doc.add_paragraph("Criterion 4: Academic Writing & Referencing (10 Marks)\nStandard academic style, typography, and citation integrity.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# =========================================================================
# TEST A: Upload and Format
# =========================================================================
def test_a_upload_and_format_preserves_original_and_creates_version():
    content = create_sample_docx()
    res = client.post(
        "/api/jobs/upload",
        files={"file": ("assignment_jane.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    job = res.json()
    job_id = job["job_id"]
    assert job["current_version_id"] == 1
    assert len(job["versions"]) == 1
    assert job["versions"][0]["operation"] == "original"
    assert job["versions"][0]["filename"] == "assignment_jane.docx"
    
    # Verify original file content is intact via download of version 1
    v1_dl = client.get(f"/api/jobs/{job_id}/versions/1/download")
    assert v1_dl.status_code == 200
    assert len(v1_dl.content) == len(content)

    # Format document
    settings = FormattingSettings()
    settings.body.font_family = "Times New Roman"
    fmt_res = client.post(f"/api/jobs/{job_id}/format", json=settings.model_dump())
    assert fmt_res.status_code == 200
    fmt_job = fmt_res.json()

    # Verify Version 2 was created
    assert fmt_job["current_version_id"] == 2
    assert len(fmt_job["versions"]) == 2
    assert fmt_job["versions"][1]["operation"] == "formatted"
    assert fmt_job["versions"][1]["source_version_id"] == 1

    # Verify Version 1 remains accessible and untouched
    v1_check = client.get(f"/api/jobs/{job_id}/versions/1/download")
    assert v1_check.status_code == 200
    assert len(v1_check.content) == len(content)

# =========================================================================
# TEST B: Formatter to Assignment Doctor Handoff
# =========================================================================
def test_b_formatter_to_doctor_handoff_analyzes_exact_version():
    content = create_sample_docx()
    job = client.post(
        "/api/jobs/upload",
        files={"file": ("report.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    job_id = job["job_id"]

    # Format to produce Version 2
    settings = FormattingSettings()
    client.post(f"/api/jobs/{job_id}/format", json=settings.model_dump())

    # Formatter hands off job_id to Doctor without re-upload
    doc_res = client.get(f"/api/doctor/diagnose/{job_id}")
    assert doc_res.status_code == 200
    diagnosis = doc_res.json()

    # Must explicitly identify analyzed version as Version 2
    assert diagnosis["version_id"] == 2
    assert "Version 2" in diagnosis["version_label"]
    assert diagnosis["is_latest_version"] is True
    assert diagnosis["overall_health_score"] > 0
    assert "citation_report" in diagnosis

# =========================================================================
# TEST C: Fix and Recheck Creates New Version
# =========================================================================
def test_c_doctor_fix_and_recheck_creates_new_version():
    content = create_sample_docx()
    job = client.post(
        "/api/jobs/upload",
        files={"file": ("draft.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    job_id = job["job_id"]

    # Format to create Version 2
    client.post(f"/api/jobs/{job_id}/format", json=FormattingSettings().model_dump())

    # Fix issues on Version 2
    fix_res = client.post(
        f"/api/doctor/fix/{job_id}",
        json={"version_id": 2, "fix_all_safe": True},
    )
    assert fix_res.status_code == 200
    fix_data = fix_res.json()

    # Updated diagnosis is for Version 3
    diag = fix_data["updated_diagnosis"]
    assert diag["version_id"] == 3
    assert "Version 3" in diag["version_label"]
    assert diag["is_latest_version"] is True

    # Check project status: has 3 versions now
    status = client.get(f"/api/jobs/{job_id}").json()
    assert len(status["versions"]) == 3
    assert status["versions"][0]["version_id"] == 1  # Original
    assert status["versions"][1]["version_id"] == 2  # Formatted
    assert status["versions"][2]["version_id"] == 3  # Repaired

    # Verify student can download any version
    v1_dl = client.get(f"/api/jobs/{job_id}/versions/1/download")
    v2_dl = client.get(f"/api/jobs/{job_id}/versions/2/download")
    v3_dl = client.get(f"/api/jobs/{job_id}/versions/3/download")
    assert v1_dl.status_code == 200
    assert v2_dl.status_code == 200
    assert v3_dl.status_code == 200

# =========================================================================
# TEST D: Rubric Reuse Across Tools
# =========================================================================
def test_d_rubric_reuse_across_doctor_and_compliance():
    content = create_sample_docx()
    job = client.post(
        "/api/jobs/upload",
        files={"file": ("report_rubric_test.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    job_id = job["job_id"]

    # Upload marking scheme once
    rubric_bytes = create_sample_rubric_docx()
    rub_res = client.post(
        f"/api/jobs/{job_id}/upload-rubric",
        files={"file": ("assessment_scheme.docx", rubric_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert rub_res.status_code == 200
    assert rub_res.json()["criteria_count"] == 4

    # Format document
    client.post(f"/api/jobs/{job_id}/format", json=FormattingSettings().model_dump())

    # Open Assignment Doctor: rubric must be automatically reused
    doc_res = client.get(f"/api/doctor/diagnose/{job_id}")
    assert doc_res.status_code == 200
    diagnosis = doc_res.json()
    assert diagnosis["rubric_used"] is True
    assert diagnosis["rubric_summary"] is not None

    # Check dedicated rubric analysis endpoint
    analysis_res = client.get(f"/api/jobs/{job_id}/rubric-analysis")
    assert analysis_res.status_code == 200
    an_data = analysis_res.json()
    assert an_data["rubric_filename"] == "assessment_scheme.docx"
    assert len(an_data["rubric_analysis"]["criteria"]) == 4

# =========================================================================
# TEST E: PowerPoint Generation from Formatted Document
# =========================================================================
def test_e_powerpoint_generation_from_project_version():
    content = create_sample_docx()
    job = client.post(
        "/api/jobs/upload",
        files={"file": ("thesis.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    job_id = job["job_id"]

    # Format to create Version 2
    settings = FormattingSettings()
    settings.cover_page.assignment_title = "Deterministic Distributed Architectures"
    settings.cover_page.student_name = "Jane Doe"
    client.post(f"/api/jobs/{job_id}/format", json=settings.model_dump())

    # Plan presentation directly from project Version 2 without re-upload
    plan_res = client.post(
        f"/api/presentation/project/{job_id}/plan",
        json={"version_id": 2, "target_slide_count": 8, "theme": "corporate_pro"},
    )
    assert plan_res.status_code == 200
    outline = plan_res.json()
    # Exactly 8 slides: Slide 0 (Title) + 7 content slides
    assert len(outline["slides"]) + 1 == 8
    assert "Distributed" in outline["title"]

    # Generate PPTX
    gen_res = client.post(
        f"/api/presentation/project/{job_id}/generate",
        json={"outline": outline, "theme": "corporate_pro", "version_id": 2},
    )
    assert gen_res.status_code == 200
    gen_data = gen_res.json()
    assert gen_data["total_slides"] == 8
    assert gen_data["source_version_id"] == 2
    assert gen_data["is_stale"] is False

    # Download and verify PPTX
    dl_res = client.get(gen_data["download_url"])
    assert dl_res.status_code == 200
    assert len(dl_res.content) > 1000

# =========================================================================
# TEST F: Version Changes & Presentation Stale Detection
# =========================================================================
def test_f_presentation_stale_warning_when_document_updated():
    content = create_sample_docx()
    job = client.post(
        "/api/jobs/upload",
        files={"file": ("project_deck.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    job_id = job["job_id"]

    # Format to create Version 2
    client.post(f"/api/jobs/{job_id}/format", json=FormattingSettings().model_dump())

    # Generate presentation from Version 2
    plan = client.post(
        f"/api/presentation/project/{job_id}/plan",
        json={"version_id": 2, "target_slide_count": 5},
    ).json()
    gen = client.post(
        f"/api/presentation/project/{job_id}/generate",
        json={"outline": plan, "theme": "premium_university", "version_id": 2},
    ).json()
    assert gen["source_version_id"] == 2
    assert gen["is_stale"] is False

    # Doctor fixes issue, generating Version 3
    client.post(f"/api/doctor/fix/{job_id}", json={"version_id": 2, "fix_all_safe": True})

    # Check presentation status: should flag is_stale = True because current_version (3) > source_version (2)
    pres_status = client.get(f"/api/presentation/project/{job_id}/status").json()
    assert pres_status["current_version_id"] == 3
    assert len(pres_status["presentations"]) == 1
    assert pres_status["presentations"][0]["is_stale"] is True
    assert pres_status["presentations"][0]["source_version_id"] == 2

# =========================================================================
# TEST G: Security and Failure Handling
# =========================================================================
def test_g_security_invalid_ids_and_safe_failure_handling():
    # 1. Invalid UUID format rejection
    bad_id = client.get("/api/jobs/not-a-valid-uuid")
    assert bad_id.status_code in [400, 422]

    # 2. Path traversal attack attempt in job_id
    traversal = client.get("/api/jobs/../../etc/passwd")
    assert traversal.status_code in [400, 404, 422]

    # 3. Non-existent job returns 404
    non_existent = client.get("/api/jobs/00000000-0000-0000-0000-000000000000")
    assert non_existent.status_code == 404

    # 4. Safe failure: formatting with broken data does not destroy original
    content = create_sample_docx()
    job = client.post(
        "/api/jobs/upload",
        files={"file": ("valid.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    job_id = job["job_id"]

    # Original version 1 is intact
    v1_orig = client.get(f"/api/jobs/{job_id}/versions/1/download")
    assert v1_orig.status_code == 200
    assert len(v1_orig.content) == len(content)

    # Corrupt settings or bad input should fail gracefully
    bad_fmt = client.post(f"/api/jobs/{job_id}/format", json={"paper_size": "INVALID_SIZE"})
    assert bad_fmt.status_code in [400, 422]

    # Original file must STILL be intact and available
    v1_check = client.get(f"/api/jobs/{job_id}/versions/1/download")
    assert v1_check.status_code == 200
    assert len(v1_check.content) == len(content)
