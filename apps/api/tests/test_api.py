import pytest
from fastapi.testclient import TestClient
from pathlib import Path
import io
import docx
from src.main import app

client = TestClient(app)

def test_health():
    res = client.get("/api/health")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "healthy"

def test_upload_and_format_workflow(tmp_path):
    # Create docx in memory
    doc = docx.Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("This is the intro.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("Ref 1.")
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    # 1. Upload
    files = {"file": ("student_work.docx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/jobs/upload", files=files)
    assert res.status_code == 200
    job = res.json()
    job_id = job["job_id"]
    assert job["inspection"]["total_paragraphs"] > 0

    # 2. Format
    format_payload = {
        "paper_size": "A4",
        "margins": {"top_inches": 1.0, "bottom_inches": 1.0, "left_inches": 1.0, "right_inches": 1.0},
        "body": {
            "font_family": "Arial",
            "font_size_pt": 11.0,
            "font_color_hex": "#111827",
            "alignment": "LEFT",
            "line_spacing": 1.5,
            "space_before_pt": 0.0,
            "space_after_pt": 6.0,
            "first_line_indent_inches": 0.0,
        },
        "heading1": {
            "font_family": "Arial",
            "font_size_pt": 16.0,
            "bold": True,
            "italic": False,
            "color_hex": "#1E3A8A",
            "alignment": "LEFT",
            "space_before_pt": 16.0,
            "space_after_pt": 6.0,
            "page_break_before": True,
        },
        "heading2": {
            "font_family": "Arial",
            "font_size_pt": 13.0,
            "bold": True,
            "italic": False,
            "color_hex": "#1E3A8A",
            "alignment": "LEFT",
            "space_before_pt": 12.0,
            "space_after_pt": 4.0,
            "page_break_before": False,
        },
        "heading3": {
            "font_family": "Arial",
            "font_size_pt": 11.0,
            "bold": True,
            "italic": True,
            "color_hex": "#374151",
            "alignment": "LEFT",
            "space_before_pt": 6.0,
            "space_after_pt": 2.0,
            "page_break_before": False,
        },
        "start_main_heading_on_new_page": True,
        "excluded_heading_break_titles": [],
        "preserve_hyperlinks": True,
        "preserve_tables": True,
        "preserve_images": True,
        "cover_page": {
            "enabled": True,
            "mode": "builtin",
            "replace_existing": True,
            "student_name": "Jane Doe",
            "student_id": "JD-12345",
            "institution": "Stanford University",
            "department": "School of Engineering",
            "course": "BSc Computer Science",
            "module_code": "CS106B",
            "assignment_title": "Programming Abstractions Project",
            "assignment_subtitle": "Final Coursework Report",
            "lecturer_name": "Prof. Smith",
            "submission_date": "October 2026",
        },
        "marking_scheme": {"enabled": False, "title": "Marking Scheme", "placement": "end_of_document", "replace_existing": False, "criteria": [], "notes": ""},
        "page_numbering": {"enabled": True, "hide_on_cover": True, "roman_front_matter": False, "restart_at_introduction": True, "footer_alignment": "RIGHT"},
        "table_of_contents": {"enabled": True, "title": "Contents", "depth": 3},
        "required_sections": ["Introduction", "References"],
    }

    res_fmt = client.post(f"/api/jobs/{job_id}/format", json=format_payload)
    assert res_fmt.status_code == 200
    fmt_data = res_fmt.json()
    assert fmt_data["is_formatted"] is True
    assert fmt_data["compliance"]["score_percentage"] >= 75

    # 3. Download
    res_dl = client.get(f"/api/jobs/{job_id}/download")
    assert res_dl.status_code == 200
    assert len(res_dl.content) > 1000

    # 4. Clean up
    res_del = client.delete(f"/api/jobs/{job_id}")
    assert res_del.status_code == 200

def test_direct_upload_alias_and_conversion_registry():
    # 1. Test /api/conversions/registry
    res_reg = client.get("/api/conversions/registry")
    assert res_reg.status_code == 200
    assert "registry" in res_reg.json()

    # 2. Test direct /api/upload alias
    doc = docx.Document()
    doc.add_heading("Direct Test", level=1)
    doc.add_paragraph("Testing /api/upload alias.")
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    files = {"file": ("direct_upload.docx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res_up = client.post("/api/upload", files=files)
    assert res_up.status_code == 200
    data = res_up.json()
    assert "job_id" in data
    assert data["original_filename"] == "direct_upload.docx"

    # Clean up
    client.delete(f"/api/jobs/{data['job_id']}")

