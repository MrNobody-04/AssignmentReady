import io
import csv
import re
from pathlib import Path
from PIL import Image
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pypdf

from src.common.errors import FileValidationError
from src.modules.presentation.ppt_generator import (
    extract_presentation_outline_from_docx,
    extract_presentation_outline_from_text,
    generate_powerpoint_presentation,
    PresentationOutline,
    SlideItem,
)
from src.modules.documents.converter import convert_docx_to_pdf, is_libreoffice_available

def sanitize_output_filename(custom_name: str, default_stem: str, target_ext: str) -> str:
    """
    Cleans and secures the user's requested output filename, ensuring
    valid characters and that the correct extension is appended.
    """
    clean_target_ext = target_ext if target_ext.startswith(".") else f".{target_ext}"
    name_to_clean = custom_name.strip() if custom_name and custom_name.strip() else default_stem

    # Strip any trailing extension if present (e.g. user entered .docx, .pdf, or anything)
    stem = Path(name_to_clean).stem if "." in name_to_clean else name_to_clean

    # Strip dangerous characters
    safe = re.sub(r'[^a-zA-Z0-9_\-\s\.]', '_', stem)
    safe = re.sub(r'\s+', ' ', safe).strip()
    if not safe or safe in ['.', '..']:
        safe = Path(default_stem).stem or "converted_document"

    return f"{safe}{clean_target_ext}"

def convert_image_to_pdf(image_bytes: bytes) -> bytes:
    """Converts image (PNG, JPG, WEBP, BMP, TIFF) to high-quality PDF document."""
    try:
        img = Image.open(io.BytesIO(image_bytes))
        # Handle alpha channel / transparency for PDF
        if img.mode in ("RGBA", "LA", "P"):
            rgb_img = Image.new("RGB", img.size, (255, 255, 255))
            if img.mode == "P":
                img = img.convert("RGBA")
            rgb_img.paste(img, mask=img.split()[-1] if img.mode == "RGBA" else None)
            img = rgb_img
        elif img.mode != "RGB":
            img = img.convert("RGB")

        out = io.BytesIO()
        img.save(out, format="PDF", resolution=100.0)
        return out.getvalue()
    except Exception as e:
        raise FileValidationError(f"Failed to convert image to PDF: {str(e)}")

def convert_image_format(image_bytes: bytes, target_fmt: str) -> bytes:
    """Converts image between standard image formats (PNG, JPEG)."""
    try:
        img = Image.open(io.BytesIO(image_bytes))
        out = io.BytesIO()
        fmt_upper = "JPEG" if target_fmt.lower() in ["jpg", "jpeg"] else "PNG"
        if fmt_upper == "JPEG" and img.mode in ("RGBA", "LA", "P"):
            rgb_img = Image.new("RGB", img.size, (255, 255, 255))
            if img.mode == "P":
                img = img.convert("RGBA")
            rgb_img.paste(img, mask=img.split()[-1] if img.mode == "RGBA" else None)
            img = rgb_img

        img.save(out, format=fmt_upper, quality=95)
        return out.getvalue()
    except Exception as e:
        raise FileValidationError(f"Failed to convert image format: {str(e)}")

def convert_csv_to_xlsx(csv_bytes: bytes) -> bytes:
    """
    Converts raw CSV file into a polished, professional Microsoft Excel workbook (.xlsx).
    Applies styled academic navy header row, zebra rows, borders, and auto column widths.
    """
    try:
        text = csv_bytes.decode("utf-8", errors="replace")
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample)
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = list(reader)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data Table"

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        body_font = Font(name="Segoe UI", size=10, color="1E293B")

        thin_side = Side(border_style="thin", color="CBD5E1")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        for row_idx, row in enumerate(rows, start=1):
            for col_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = border
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = body_font
                    if row_idx % 2 == 0:
                        cell.fill = zebra_fill
                    cell.alignment = Alignment(vertical="center")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Freeze header row
        ws.freeze_panes = "A2"

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    except Exception as e:
        raise FileValidationError(f"Failed to convert CSV to XLSX: {str(e)}")

def convert_txt_to_docx(text_content: str) -> bytes:
    """Converts plain text into a properly formatted Microsoft Word document."""
    doc = docx.Document()

    # Set 1-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Base style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    lines = text_content.splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("# "):
            h = doc.add_heading(stripped.lstrip("#").strip(), level=1)
            h.style.font.name = 'Times New Roman'
            h.style.font.size = Pt(16)
            h.style.font.bold = True
        elif stripped.startswith("## "):
            h = doc.add_heading(stripped.lstrip("#").strip(), level=2)
            h.style.font.name = 'Times New Roman'
            h.style.font.size = Pt(14)
            h.style.font.bold = True
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(stripped[2:].strip(), style='List Bullet')
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph(stripped)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

def convert_pdf_to_docx(pdf_bytes: bytes) -> bytes:
    """
    Extracts text from PDF and structures it into an editable Word document (.docx).
    Preserves multi-page flow, headers, and bullet paragraphs.
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        doc = docx.Document()

        for sec in doc.sections:
            sec.top_margin = Inches(1.0)
            sec.bottom_margin = Inches(1.0)
            sec.left_margin = Inches(1.0)
            sec.right_margin = Inches(1.0)

        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            lines = [l.strip() for l in text.splitlines() if l.strip()]

            if page_idx > 0:
                doc.add_page_break()

            for line in lines:
                # Detect heading
                if len(line) < 60 and (line.isupper() or re.match(r"^\d+\.\s+[A-Z]", line)):
                    h = doc.add_heading(line, level=1)
                    h.style.font.name = 'Times New Roman'
                    h.style.font.size = Pt(16)
                    h.style.font.bold = True
                elif line.startswith("•") or line.startswith("-"):
                    p = doc.add_paragraph(line.lstrip("•-").strip(), style='List Bullet')
                    p.paragraph_format.line_spacing = 1.5
                else:
                    p = doc.add_paragraph(line)
                    p.paragraph_format.line_spacing = 1.5
                    p.paragraph_format.space_after = Pt(6)

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception as e:
        raise FileValidationError(f"Failed to convert PDF to DOCX: {str(e)}")

def convert_csv_to_docx_table(csv_bytes: bytes) -> bytes:
    """Converts CSV into an academic Word document table."""
    try:
        text = csv_bytes.decode("utf-8", errors="replace")
        sample = text[:2048]
        try:
            delimiter = csv.Sniffer().sniff(sample).delimiter
        except Exception:
            delimiter = ","

        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        doc = docx.Document()

        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Table Grid'
            for r_idx, row in enumerate(rows):
                for c_idx, val in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = str(val)
                    if r_idx == 0:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.bold = True

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception as e:
        raise FileValidationError(f"Failed to convert CSV to Word Table: {str(e)}")

def convert_source_to_pptx(
    source_bytes: bytes,
    source_format: str,
    theme: str = "academic_classic",
) -> bytes:
    """Extracts presentation outline and generates PowerPoint presentation (.pptx)."""
    if source_format == "docx":
        # Load docx from temp stream
        doc_stream = io.BytesIO(source_bytes)
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(source_bytes)
            tmp_path = Path(tmp.name)
        try:
            outline = extract_presentation_outline_from_docx(tmp_path)
        finally:
            if tmp_path.exists():
                tmp_path.unlink()
    elif source_format == "pdf":
        reader = pypdf.PdfReader(io.BytesIO(source_bytes))
        full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
        outline = extract_presentation_outline_from_text(full_text)
    elif source_format in ["txt", "csv"]:
        text = source_bytes.decode("utf-8", errors="replace")
        outline = extract_presentation_outline_from_text(text)
    else:
        raise FileValidationError(f"Presentation generation not supported for format '{source_format}'.")

    return generate_powerpoint_presentation(outline, theme_key=theme)
