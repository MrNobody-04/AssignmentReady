import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from src.modules.jobs.job_schema import (
    FormattingSettings,
    InspectionResult,
    CaptionItem,
    CoverPageData,
    DeclarationSettings,
    AcknowledgementSettings,
    AbstractSettings,
    ListOfFiguresSettings,
    ListOfTablesSettings,
    ListOfAbbreviationsSettings,
)

TOC_MARKERS = ["table of contents", "contents", "table of content"]
LOF_MARKERS = ["list of figures", "table of figures", "figure list"]
LOT_MARKERS = ["list of tables", "table of tables", "table list"]
DECLARATION_MARKERS = ["declaration", "statement of originality", "declaration of originality", "author's declaration"]
ACKNOWLEDGEMENT_MARKERS = ["acknowledgement", "acknowledgments", "acknowledging"]
ABSTRACT_MARKERS = ["abstract", "executive summary"]
ABBREVIATION_MARKERS = ["list of abbreviations", "abbreviations", "acronyms"]

def _create_field_code(instr_text: str, fallback_text: str = "") -> OxmlElement:
    """Creates a native Word field code element (fldSimple)."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr_text)
    if fallback_text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = fallback_text
        r.append(t)
        fld.append(r)
    return fld

def add_section_break_to_paragraph(p) -> None:
    """Inserts a NextPage section break into the paragraph's pPr."""
    pPr = p._p.get_or_add_pPr()
    sectPr = OxmlElement("w:sectPr")
    type_elem = OxmlElement("w:type")
    type_elem.set(qn("w:val"), "nextPage")
    sectPr.append(type_elem)
    pPr.append(sectPr)

def remove_existing_front_matter_paragraphs(doc: docx.Document, settings: FormattingSettings) -> None:
    """
    Detects and removes existing duplicate front matter paragraphs (old cover page,
    old TOC, old LOF, old LOT, old declaration) so they do not collide with newly generated front matter.
    """
    paragraphs_to_remove = []
    
    # 1. Old cover page removal if replace_existing is True
    if settings.cover_page.enabled and settings.cover_page.replace_existing:
        cover_kws = ["student name", "student id", "matriculation", "course code", "submitted by", "submitted to", "lecturer"]
        cover_end_idx = -1
        for idx, p in enumerate(doc.paragraphs[:20]):
            txt = p.text.strip().lower()
            if any(kw in txt for kw in cover_kws):
                cover_end_idx = idx
            if p.paragraph_format.page_break_before and cover_end_idx >= 0:
                break
        if cover_end_idx >= 0:
            for i in range(cover_end_idx + 1):
                paragraphs_to_remove.append(doc.paragraphs[i])

    # 2. Old TOC, LOF, LOT, Declaration removal
    collecting_removal = False
    for p in doc.paragraphs:
        txt = p.text.strip().lower()
        if not txt:
            if collecting_removal:
                paragraphs_to_remove.append(p)
            continue

        is_front_matter_heading = False
        if settings.table_of_contents.enabled and any(txt == kw or txt.startswith(kw + ":") for kw in TOC_MARKERS):
            is_front_matter_heading = True
        elif settings.list_of_figures.enabled and any(txt == kw or txt.startswith(kw + ":") for kw in LOF_MARKERS):
            is_front_matter_heading = True
        elif settings.list_of_tables.enabled and any(txt == kw or txt.startswith(kw + ":") for kw in LOT_MARKERS):
            is_front_matter_heading = True
        elif settings.declaration.enabled and any(txt == kw or txt.startswith(kw + ":") for kw in DECLARATION_MARKERS):
            is_front_matter_heading = True
        elif settings.acknowledgement.enabled and any(txt == kw or txt.startswith(kw + ":") for kw in ACKNOWLEDGEMENT_MARKERS):
            is_front_matter_heading = True
        elif settings.abstract.enabled and any(txt == kw or txt.startswith(kw + ":") for kw in ABSTRACT_MARKERS):
            is_front_matter_heading = True

        if is_front_matter_heading:
            collecting_removal = True
            paragraphs_to_remove.append(p)
            continue

        if collecting_removal:
            style_name = p.style.name.lower() if p.style else ""
            is_next_real_heading = (
                "heading 1" in style_name
                or "heading 2" in style_name
                or re.match(r"^(?:chapter\s+\d+|1\.\d+|introduction)", txt)
            )
            if is_next_real_heading:
                collecting_removal = False
            else:
                paragraphs_to_remove.append(p)

    for p in set(paragraphs_to_remove):
        if p._p.getparent() is not None:
            p._p.getparent().remove(p._p)

def insert_declaration_before(target_p, settings: DeclarationSettings, base_font: str = "Times New Roman") -> None:
    """Inserts Declaration of Originality page before target_p."""
    p_title = target_p.insert_paragraph_before(settings.title, style="Heading 1")
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(18)

    p_body = target_p.insert_paragraph_before(settings.statement)
    p_body.paragraph_format.line_spacing = 1.5
    p_body.paragraph_format.space_after = Pt(24)
    for r in p_body.runs:
        r.font.name = base_font
        r.font.size = Pt(12)

    if settings.include_signature_line:
        p_sig = target_p.insert_paragraph_before()
        p_sig.paragraph_format.space_before = Pt(36)
        p_sig.paragraph_format.space_after = Pt(6)
        r_sig = p_sig.add_run("Student Signature: ___________________________________     Date: ______________")
        r_sig.font.name = base_font
        r_sig.font.size = Pt(11)

    p_break = target_p.insert_paragraph_before()
    p_break.paragraph_format.page_break_before = True

def insert_acknowledgement_before(target_p, settings: AcknowledgementSettings, base_font: str = "Times New Roman") -> None:
    """Inserts Acknowledgements page before target_p."""
    p_title = target_p.insert_paragraph_before(settings.title, style="Heading 1")
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(18)

    content = settings.content.strip() or (
        "I would like to express my sincere appreciation to my lecturer, academic supervisors, "
        "and institution faculty for their guidance, valuable insights, and continuous support "
        "throughout the completion of this assignment."
    )
    p_body = target_p.insert_paragraph_before(content)
    p_body.paragraph_format.line_spacing = 1.5
    p_body.paragraph_format.space_after = Pt(24)
    for r in p_body.runs:
        r.font.name = base_font
        r.font.size = Pt(12)

    p_break = target_p.insert_paragraph_before()
    p_break.paragraph_format.page_break_before = True

def insert_abstract_before(target_p, settings: AbstractSettings, base_font: str = "Times New Roman") -> None:
    """Inserts Abstract / Executive Summary page before target_p."""
    p_title = target_p.insert_paragraph_before(settings.title, style="Heading 1")
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(18)

    content = settings.content.strip() or (
        "This academic report provides an in-depth examination and structured analysis of the subject matter. "
        "It outlines the core research questions, theoretical context, investigative methodology, and synthesized conclusions."
    )
    p_body = target_p.insert_paragraph_before(content)
    p_body.paragraph_format.line_spacing = 1.5
    p_body.paragraph_format.space_after = Pt(24)
    for r in p_body.runs:
        r.font.name = base_font
        r.font.size = Pt(12)

    p_break = target_p.insert_paragraph_before()
    p_break.paragraph_format.page_break_before = True

def insert_table_of_contents_before(
    target_p,
    headings: list,
    title: str = "Table of Contents",
    base_font: str = "Times New Roman"
) -> None:
    """
    Inserts Table of Contents section before target_p containing both native Word TOC field code
    and pre-rendered tabbed entries with leader dots.
    """
    p_title = target_p.insert_paragraph_before(title, style="Heading 1")
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(12)

    # Word native field code paragraph
    p_fld = target_p.insert_paragraph_before()
    p_fld.paragraph_format.space_after = Pt(6)
    run_fld = p_fld.add_run()
    run_fld._r.append(_create_field_code('TOC \\o "1-3" \\h \\z \\u', "Table of Contents (Auto-updating field)"))
    run_fld.font.size = Pt(9)
    run_fld.font.italic = True
    run_fld.font.color.rgb = RGBColor(100, 116, 139)

    # Pre-rendered entries for instant visual fidelity
    if headings:
        skip_titles = set(TOC_MARKERS + LOF_MARKERS + LOT_MARKERS + DECLARATION_MARKERS + ACKNOWLEDGEMENT_MARKERS + ABSTRACT_MARKERS)
        est_page = 1
        for h in headings:
            h_text = h.text if hasattr(h, "text") else h.get("text", "")
            h_level = h.level if hasattr(h, "level") else h.get("level", 1)
            if not h_text or h_text.lower() in skip_titles:
                continue

            entry_p = target_p.insert_paragraph_before()
            entry_p.paragraph_format.space_before = Pt(2)
            entry_p.paragraph_format.space_after = Pt(2)
            entry_p.paragraph_format.line_spacing = 1.15

            indent = (h_level - 1) * 0.3
            if indent > 0:
                entry_p.paragraph_format.left_indent = Inches(indent)

            tab_stops = entry_p.paragraph_format.tab_stops
            tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

            r_text = entry_p.add_run(h_text)
            r_text.font.name = base_font
            r_text.font.size = Pt(11 if h_level == 1 else 10.5)
            if h_level == 1:
                r_text.font.bold = True

            r_tab = entry_p.add_run("\t")
            r_page = entry_p.add_run(str(est_page))
            r_page.font.name = base_font
            r_page.font.size = Pt(10)
            if h_level == 1:
                est_page += 1

    p_break = target_p.insert_paragraph_before()
    p_break.paragraph_format.page_break_before = True

def insert_list_of_figures_before(
    target_p,
    figures: list[CaptionItem],
    title: str = "List of Figures",
    base_font: str = "Times New Roman"
) -> None:
    """Inserts List of Figures page before target_p."""
    p_title = target_p.insert_paragraph_before(title, style="Heading 1")
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(12)

    # Word native field code
    p_fld = target_p.insert_paragraph_before()
    p_fld.paragraph_format.space_after = Pt(6)
    run_fld = p_fld.add_run()
    run_fld._r.append(_create_field_code('TOC \\c "Figure"', "List of Figures (Auto-updating field)"))
    run_fld.font.size = Pt(9)
    run_fld.font.italic = True
    run_fld.font.color.rgb = RGBColor(100, 116, 139)

    # Pre-rendered entries with tab stops and dot leaders
    est_page = 1
    for fig in figures:
        p_item = target_p.insert_paragraph_before()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15

        tab_stops = p_item.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

        label_run = p_item.add_run(f"Figure {fig.number}: ")
        label_run.font.name = base_font
        label_run.font.size = Pt(10.5)
        label_run.font.bold = True

        cap_run = p_item.add_run(fig.caption)
        cap_run.font.name = base_font
        cap_run.font.size = Pt(10.5)

        p_item.add_run("\t")
        pg_run = p_item.add_run(str(est_page))
        pg_run.font.name = base_font
        pg_run.font.size = Pt(10)
        est_page += 1

    p_break = target_p.insert_paragraph_before()
    p_break.paragraph_format.page_break_before = True

def insert_list_of_tables_before(
    target_p,
    tables: list[CaptionItem],
    title: str = "List of Tables",
    base_font: str = "Times New Roman"
) -> None:
    """Inserts List of Tables page before target_p."""
    p_title = target_p.insert_paragraph_before(title, style="Heading 1")
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(12)

    # Word native field code
    p_fld = target_p.insert_paragraph_before()
    p_fld.paragraph_format.space_after = Pt(6)
    run_fld = p_fld.add_run()
    run_fld._r.append(_create_field_code('TOC \\c "Table"', "List of Tables (Auto-updating field)"))
    run_fld.font.size = Pt(9)
    run_fld.font.italic = True
    run_fld.font.color.rgb = RGBColor(100, 116, 139)

    # Pre-rendered entries with tab stops and dot leaders
    est_page = 1
    for tbl in tables:
        p_item = target_p.insert_paragraph_before()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15

        tab_stops = p_item.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

        label_run = p_item.add_run(f"Table {tbl.number}: ")
        label_run.font.name = base_font
        label_run.font.size = Pt(10.5)
        label_run.font.bold = True

        cap_run = p_item.add_run(tbl.caption)
        cap_run.font.name = base_font
        cap_run.font.size = Pt(10.5)

        p_item.add_run("\t")
        pg_run = p_item.add_run(str(est_page))
        pg_run.font.name = base_font
        pg_run.font.size = Pt(10)
        est_page += 1

    p_break = target_p.insert_paragraph_before()
    p_break.paragraph_format.page_break_before = True

def insert_list_of_abbreviations_before(
    target_p,
    settings: ListOfAbbreviationsSettings,
    base_font: str = "Times New Roman"
) -> None:
    """Inserts List of Abbreviations page before target_p."""
    p_title = target_p.insert_paragraph_before(settings.title, style="Heading 1")
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(18)

    items = settings.items
    if not items:
        items = [
            type("Abbr", (), {"term": "API", "definition": "Application Programming Interface"})(),
            type("Abbr", (), {"term": "UI / UX", "definition": "User Interface / User Experience"})(),
            type("Abbr", (), {"term": "DOCX", "definition": "Office Open XML Document"})(),
            type("Abbr", (), {"term": "PDF", "definition": "Portable Document Format"})(),
        ]

    for item in items:
        p_item = target_p.insert_paragraph_before()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15

        tab_stops = p_item.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(2.0), WD_TAB_ALIGNMENT.LEFT)

        r_term = p_item.add_run(f"{item.term:<12}")
        r_term.font.name = base_font
        r_term.font.size = Pt(10.5)
        r_term.font.bold = True

        p_item.add_run("\t")

        r_def = p_item.add_run(item.definition)
        r_def.font.name = base_font
        r_def.font.size = Pt(10.5)

    p_break = target_p.insert_paragraph_before()
    p_break.paragraph_format.page_break_before = True
