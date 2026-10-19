import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt
from src.modules.jobs.job_schema import InspectionResult, DocumentStructureItem
from src.common.errors import DocumentProcessingError

COVER_KEYWORDS = [
    "student name", "student id", "matriculation", "course code",
    "module code", "submitted by", "submitted to", "lecturer",
    "instructor", "assignment title", "date of submission", "academic year"
]

TOC_KEYWORDS = ["table of contents", "contents"]
REF_KEYWORDS = ["references", "bibliography", "works cited"]
RUBRIC_KEYWORDS = ["marking scheme", "marking criteria", "assessment criteria", "grading rubric", "rubric"]

def inspect_docx(file_path: Path) -> InspectionResult:
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        raise DocumentProcessingError(f"Failed to parse Word document: {str(e)}")

    total_paragraphs = len(doc.paragraphs)
    total_words = 0
    detected_headings: list[DocumentStructureItem] = []
    first_body_font: str | None = None
    first_body_size: float | None = None

    has_cover_markers = False
    has_toc = False
    has_refs = False
    has_rubric = False

    # Check sections for margins
    margins_inches: dict[str, float] = {}
    if doc.sections:
        first_sec = doc.sections[0]
        margins_inches = {
            "top": round(first_sec.top_margin.inches, 2) if first_sec.top_margin else 1.0,
            "bottom": round(first_sec.bottom_margin.inches, 2) if first_sec.bottom_margin else 1.0,
            "left": round(first_sec.left_margin.inches, 2) if first_sec.left_margin else 1.0,
            "right": round(first_sec.right_margin.inches, 2) if first_sec.right_margin else 1.0,
        }

    # Inspect paragraphs
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        words = len(text.split())
        total_words += words

        lower_text = text.lower()
        style_name = p.style.name if p.style else ""

        # Cover page markers check in first 10 paragraphs
        if idx < 10:
            if any(kw in lower_text for kw in COVER_KEYWORDS):
                has_cover_markers = True

        # Heading detection
        level = 0
        if style_name.startswith("Heading 1") or style_name.endswith("1") and "Heading" in style_name:
            level = 1
        elif style_name.startswith("Heading 2") or style_name.endswith("2") and "Heading" in style_name:
            level = 2
        elif style_name.startswith("Heading 3") or style_name.endswith("3") and "Heading" in style_name:
            level = 3
        elif style_name.lower() in ["title"]:
            level = 1
        elif text and len(text) < 80 and not text.endswith(".") and words > 0:
            # Check if all runs are bold
            if p.runs and all(r.bold for r in p.runs if r.text.strip()):
                level = 1

        if level > 0 and text:
            detected_headings.append(
                DocumentStructureItem(level=level, text=text, word_count=words)
            )

            # Check specific sections
            if any(kw in lower_text for kw in TOC_KEYWORDS):
                has_toc = True
            if any(kw in lower_text for kw in REF_KEYWORDS):
                has_refs = True
            if any(kw in lower_text for kw in RUBRIC_KEYWORDS):
                has_rubric = True

    # Determine dominant body font by character count
    font_char_counts: dict[str, int] = {}
    for p in doc.paragraphs:
        s_name = p.style.name.lower() if p.style else ""
        if "heading" not in s_name and "title" not in s_name:
            for r in p.runs:
                f_name = r.font.name
                if f_name and r.text:
                    font_char_counts[f_name] = font_char_counts.get(f_name, 0) + len(r.text.strip())

    if font_char_counts:
        first_body_font = max(font_char_counts, key=font_char_counts.get)

    # Count tables
    total_tables = len(doc.tables)

    # Count images by searching OOXML for drawing/blip elements
    total_images = 0
    try:
        xml_content = doc._element.xml
        total_images = len(re.findall(r"<a:blip|<w:drawing", xml_content))
    except Exception:
        total_images = 0

    # Also check if TOC XML field exists
    if not has_toc:
        try:
            if "TOC " in doc._element.xml:
                has_toc = True
        except Exception:
            pass

    return InspectionResult(
        filename=file_path.name,
        file_size_bytes=file_path.stat().st_size if file_path.exists() else 0,
        total_paragraphs=total_paragraphs,
        total_words=total_words,
        total_tables=total_tables,
        total_images=total_images,
        detected_headings=detected_headings,
        detected_body_font=first_body_font or "Times New Roman",
        detected_margins_inches=margins_inches,
        has_cover_page_markers=has_cover_markers,
        has_table_of_contents=has_toc,
        has_references_section=has_refs,
        has_marking_scheme_section=has_rubric,
    )
