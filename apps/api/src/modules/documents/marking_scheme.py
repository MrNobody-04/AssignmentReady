import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from src.modules.jobs.job_schema import MarkingSchemeData

def _set_cell_background(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def _set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for edge, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def apply_marking_scheme(
    doc: docx.Document,
    scheme_data: MarkingSchemeData,
    settings=None,
    uploaded_file_path=None,
    anchor_paragraph=None,
) -> None:
    if not scheme_data.enabled:
        return

    criteria = scheme_data.criteria
    if not criteria:
        # If no criteria explicitly passed, provide standard academic criteria default
        from src.modules.jobs.job_schema import MarkingCriterion
        criteria = [
            MarkingCriterion(name="Understanding & Core Concepts", description="Demonstrates thorough grasp of theoretical domain knowledge and context.", max_marks=30.0, weight_percentage=30.0),
            MarkingCriterion(name="Analysis & Critical Evaluation", description="Evidence of critical thought, rigorous evaluation, and sound methodology.", max_marks=30.0, weight_percentage=30.0),
            MarkingCriterion(name="Structure & Coherence", description="Logical flow, clear section hierarchy, professional academic terminology.", max_marks=20.0, weight_percentage=20.0),
            MarkingCriterion(name="Referencing & Academic Integrity", description="Accurate in-text citations and complete bibliography adhering to required style.", max_marks=20.0, weight_percentage=20.0),
        ]

    # Find placement
    insert_before_p = anchor_paragraph
    if insert_before_p is None and scheme_data.placement == "before_references":
        for p in doc.paragraphs:
            if p.text.strip().lower() in ["references", "bibliography", "works cited"]:
                insert_before_p = p
                break
    elif insert_before_p is None and scheme_data.placement == "after_cover":
        # First paragraph after cover page
        for p in doc.paragraphs:
            if p.paragraph_format.page_break_before:
                insert_before_p = p
                break

    # Add section title
    if insert_before_p:
        h = insert_before_p.insert_paragraph_before(scheme_data.title, style="Heading 1")
    else:
        h = doc.add_paragraph(scheme_data.title, style="Heading 1")
    h.paragraph_format.page_break_before = True

    if scheme_data.notes:
        if insert_before_p:
            p_note = insert_before_p.insert_paragraph_before(scheme_data.notes)
        else:
            p_note = doc.add_paragraph(scheme_data.notes)
        p_note.paragraph_format.space_after = Pt(8)
        if p_note.runs:
            p_note.runs[0].font.size = Pt(10)
            p_note.runs[0].font.italic = True
            p_note.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    # Insert Table: columns: [Criterion, Description, Max Marks, Weight]
    # Note: python-docx tables are added to doc. To place before a paragraph, we insert it via XML or add at end
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Column widths
    col_widths = [Inches(1.8), Inches(3.2), Inches(0.9), Inches(0.9)]

    # Header Row
    hdr_cells = table.rows[0].cells
    headers = ["Criterion", "Assessment Focus / Descriptor", "Max Marks", "Weight"]
    for i, header_text in enumerate(headers):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        cell.text = header_text
        _set_cell_background(cell, "1E3A8A")  # Deep Navy
        _set_cell_margins(cell, 140, 140, 160, 160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    total_marks = 0.0
    total_weight = 0.0

    # Data Rows
    for idx, c in enumerate(criteria):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        total_marks += c.max_marks
        total_weight += c.weight_percentage

        values = [c.name, c.description, f"{c.max_marks:.0f}", f"{c.weight_percentage:.0f}%" if c.weight_percentage else "-"]
        for i, val in enumerate(values):
            cell = row_cells[i]
            cell.width = col_widths[i]
            cell.text = val
            _set_cell_background(cell, bg_color)
            _set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9.0)
                if i == 0:
                    r.font.bold = True
                r.font.color.rgb = RGBColor(30, 41, 59)

    # Summary Row
    summary_cells = table.add_row().cells
    summary_cells[0].text = "Total"
    summary_cells[1].text = "Aggregate Module Assessment"
    summary_cells[2].text = f"{total_marks:.0f}"
    summary_cells[3].text = f"{total_weight:.0f}%" if total_weight > 0 else "100%"

    for i in range(4):
        cell = summary_cells[i]
        cell.width = col_widths[i]
        _set_cell_background(cell, "E2E8F0")
        _set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)

    # If user selected placement other than end, move table XML before target paragraph
    if insert_before_p:
        insert_before_p._p.addprevious(h._p)
        insert_before_p._p.addprevious(table._tbl)
