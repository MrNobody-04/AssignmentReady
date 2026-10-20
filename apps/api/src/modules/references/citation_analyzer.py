import re
from pathlib import Path
from pydantic import BaseModel, Field
import docx

class InTextCitation(BaseModel):
    raw_text: str
    author: str
    year: str
    paragraph_index: int
    citation_type: str  # "author_date", "numeric"

class ReferenceEntry(BaseModel):
    entry_index: int
    raw_text: str
    detected_author: str = ""
    detected_year: str = ""
    is_complete: bool = True
    missing_fields: list[str] = Field(default_factory=list)

class CitationAnalysisReport(BaseModel):
    total_in_text_citations: int = 0
    total_references_entries: int = 0
    detected_style: str = "Author-Date (Harvard/APA)"  # "Author-Date", "Numeric (IEEE)", "Mixed / Inconsistent"
    unmatched_in_text_citations: list[InTextCitation] = Field(default_factory=list)
    uncited_reference_entries: list[ReferenceEntry] = Field(default_factory=list)
    style_inconsistencies: list[str] = Field(default_factory=list)
    incomplete_reference_entries: list[ReferenceEntry] = Field(default_factory=list)
    is_consistent: bool = True
    summary: str = ""

# Regex patterns for in-text citations
AUTHOR_DATE_PATTERN = re.compile(
    r"\(([A-Z][a-zA-Z\s\.\-&]+(?:et\s+al\.?)?),\s*(\d{4}[a-z]?)(?:,\s*p{1,2}\.?\s*\d+(?:-\d+)?)?\)",
)
NUMERIC_PATTERN = re.compile(r"\[(\d+(?:\s*,\s*\d+|\s*-\s*\d+)*)\]")

def analyze_document_citations(doc_path: Path) -> CitationAnalysisReport:
    """
    Deterministically analyzes in-text citations and reference list entries in a Word document.
    Identifies unmatched citations, uncited bibliography entries, and style inconsistencies.
    """
    try:
        doc = docx.Document(str(doc_path))
    except Exception as e:
        return CitationAnalysisReport(summary=f"Could not open document: {str(e)}")

    in_text_citations: list[InTextCitation] = []
    reference_entries: list[ReferenceEntry] = []
    has_author_date = False
    has_numeric = False

    # 1. Locate References Section
    ref_heading_idx = -1
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip().lower()
        if text in ["references", "bibliography", "works cited", "reference list"]:
            ref_heading_idx = idx
            break

    # If no explicit heading found, check for heading style with 'reference'
    if ref_heading_idx == -1:
        for idx, p in enumerate(doc.paragraphs):
            style_name = (p.style.name if p.style else "").lower()
            text = p.text.strip().lower()
            if "heading" in style_name and any(k in text for k in ["references", "bibliography"]):
                ref_heading_idx = idx
                break

    # 2. Extract In-Text Citations from Body Paragraphs
    body_paragraphs = doc.paragraphs[:ref_heading_idx] if ref_heading_idx != -1 else doc.paragraphs

    for idx, p in enumerate(body_paragraphs):
        text = p.text.strip()
        if not text:
            continue

        # Look for Author-Date: (Smith, 2021) or (Smith & Jones, 2020)
        for m in AUTHOR_DATE_PATTERN.finditer(text):
            has_author_date = True
            author_str = m.group(1).strip()
            year_str = m.group(2).strip()
            in_text_citations.append(
                InTextCitation(
                    raw_text=m.group(0),
                    author=author_str,
                    year=year_str,
                    paragraph_index=idx + 1,
                    citation_type="author_date",
                )
            )

        # Look for Numeric: [1], [2, 3]
        for m in NUMERIC_PATTERN.finditer(text):
            has_numeric = True
            nums = m.group(1).strip()
            in_text_citations.append(
                InTextCitation(
                    raw_text=m.group(0),
                    author="",
                    year=nums,
                    paragraph_index=idx + 1,
                    citation_type="numeric",
                )
            )

    # 3. Extract Reference List Entries
    if ref_heading_idx != -1:
        ref_p_idx = 1
        for p in doc.paragraphs[ref_heading_idx + 1:]:
            text = p.text.strip()
            if not text:
                continue

            # Check if this looks like another major section (e.g., Appendix)
            style_name = (p.style.name if p.style else "").lower()
            if "heading 1" in style_name and "appendix" in text.lower():
                break

            missing: list[str] = []
            # Extract year if present
            year_match = re.search(r"\((\d{4}[a-z]?)\)|(?:,\s*(\d{4}[a-z]?)\.)", text)
            detected_year = year_match.group(1) or year_match.group(2) if year_match else ""
            if not detected_year and len(text) > 30:
                missing.append("Publication Year")

            # Extract author (typically first token before year or first comma)
            author_match = re.match(r"^([A-Z][a-zA-Z\s\.\-]+?)(?:,\s|\s\()", text)
            detected_author = author_match.group(1).strip() if author_match else text.split()[0] if text.split() else ""
            if not detected_author:
                missing.append("Author Name")

            is_comp = len(missing) == 0

            reference_entries.append(
                ReferenceEntry(
                    entry_index=ref_p_idx,
                    raw_text=text,
                    detected_author=detected_author,
                    detected_year=detected_year,
                    is_complete=is_comp,
                    missing_fields=missing,
                )
            )
            ref_p_idx += 1

    # 4. Cross-Reference Matching
    unmatched_in_text: list[InTextCitation] = []
    uncited_refs: list[ReferenceEntry] = []

    # Check in-text citations against references list
    ref_texts = [r.raw_text.lower() for r in reference_entries]
    for cit in in_text_citations:
        if cit.citation_type == "author_date":
            # Match author surname and year in any reference entry
            first_author = cit.author.split()[0].replace(",", "").lower()
            matched = any(first_author in r_text and cit.year in r_text for r_text in ref_texts)
            if not matched:
                unmatched_in_text.append(cit)

    # Check references against in-text citations
    all_body_text = " ".join(p.text for p in body_paragraphs).lower()
    for ref in reference_entries:
        cited = False
        if ref.detected_author and len(ref.detected_author) > 2:
            author_surname = ref.detected_author.split()[0].replace(",", "").lower()
            if author_surname in all_body_text:
                cited = True
        elif ref.detected_year and ref.detected_year in all_body_text:
            cited = True

        if not cited and len(ref.raw_text) > 15:
            uncited_refs.append(ref)

    # 5. Style Consistency
    style_inconsistencies: list[str] = []
    if has_author_date and has_numeric:
        detected_style = "Mixed / Inconsistent"
        style_inconsistencies.append(
            "Document contains both Author-Date (e.g. (Smith, 2021)) and Numeric (e.g. [1]) citations. Academic conventions require a single uniform citation system throughout."
        )
    elif has_numeric:
        detected_style = "Numeric (IEEE / Nature)"
    else:
        detected_style = "Author-Date (Harvard / APA)"

    incomplete_refs = [r for r in reference_entries if not r.is_complete]

    is_consistent = (
        len(unmatched_in_text) == 0
        and len(style_inconsistencies) == 0
        and len(incomplete_refs) == 0
    )

    summary_parts = []
    summary_parts.append(f"Identified {len(in_text_citations)} in-text citation(s) and {len(reference_entries)} reference list entry/entries using {detected_style} style.")
    if unmatched_in_text:
        summary_parts.append(f"{len(unmatched_in_text)} in-text citation(s) appear to lack matching entries in the reference list.")
    if uncited_refs:
        summary_parts.append(f"{len(uncited_refs)} reference(s) appear in the bibliography without corresponding in-text citations.")
    if incomplete_refs:
        summary_parts.append(f"{len(incomplete_refs)} reference entry/entries may have missing publication year or author details.")

    return CitationAnalysisReport(
        total_in_text_citations=len(in_text_citations),
        total_references_entries=len(reference_entries),
        detected_style=detected_style,
        unmatched_in_text_citations=unmatched_in_text,
        uncited_reference_entries=uncited_refs,
        style_inconsistencies=style_inconsistencies,
        incomplete_reference_entries=incomplete_refs,
        is_consistent=is_consistent,
        summary=" ".join(summary_parts),
    )
