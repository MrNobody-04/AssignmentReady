"""Generates complete suite of test fixtures required for rigorous automated testing."""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls
import pymupdf


def make_dummy_logo() -> bytes:
    """Creates a simple 100x100 PNG logo in memory."""
    doc = pymupdf.open()
    page = doc.new_page(width=100, height=100)
    # Draw blue university crest circle
    page.draw_circle(pymupdf.Point(50, 50), 40, color=(0.1, 0.2, 0.5), fill=(0.15, 0.3, 0.7))
    pix = page.get_pixmap()
    b = pix.tobytes("png")
    doc.close()
    return b


def create_fixtures(base_dir: Path):
    base_dir.mkdir(parents=True, exist_ok=True)
    logo_bytes = make_dummy_logo()
    logo_path = base_dir / "temp_logo.png"
    with open(logo_path, "wb") as f:
        f.write(logo_bytes)

    # -------------------------------------------------------------
    # 1. Cover with Logo and Table (.docx)
    # -------------------------------------------------------------
    doc1 = docx.Document()
    p_logo = doc1.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(30)
    p_logo.add_run().add_picture(str(logo_path), width=Inches(1.5))

    p_title = doc1.add_paragraph("FACULTY OF ADVANCED COMPUTING & INFORMATICS")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.runs[0].font.name = "Arial"
    p_title.runs[0].font.size = Pt(16)
    p_title.runs[0].font.bold = True
    p_title.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc1.add_paragraph("Official Formative Coursework Assessment")
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.runs[0].font.size = Pt(12)
    p_sub.paragraph_format.space_after = Pt(40)

    tbl = doc1.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ("Student Name:", "John Alexander Smith"),
        ("Student ID:", "STU-8892014"),
        ("Module Title:", "Deep Neural Systems (CS702)"),
        ("Submission Date:", "October 24, 2026"),
    ]
    for r_idx, (k, v) in enumerate(data):
        row = tbl.rows[r_idx]
        row.cells[0].text = k
        row.cells[1].text = v
        for c in row.cells:
            c.paragraphs[0].runs[0].font.size = Pt(10)
            c.paragraphs[0].runs[0].font.name = "Arial"
            tcPr = c._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
    doc1.save(str(base_dir / "cover_with_logo_and_table.docx"))

    # -------------------------------------------------------------
    # 2. Cover with Custom Font, Header & Footer (.docx)
    # -------------------------------------------------------------
    doc2 = docx.Document()
    sec2 = doc2.sections[0]
    sec2.header.paragraphs[0].text = "UNIVERSITY OF CAMBRIDGE — DEPARTMENT OF COMPUTER SCIENCE"
    sec2.header.paragraphs[0].runs[0].font.name = "Georgia"
    sec2.header.paragraphs[0].runs[0].font.size = Pt(8.5)
    sec2.footer.paragraphs[0].text = "CONFIDENTIAL CANDIDATE SUBMISSION — DO NOT REDISTRIBUTE"
    sec2.footer.paragraphs[0].runs[0].font.name = "Georgia"
    sec2.footer.paragraphs[0].runs[0].font.size = Pt(8.5)

    p_cov2 = doc2.add_paragraph("RESEARCH DISSERTATION ARCHIVE")
    p_cov2.paragraph_format.space_before = Pt(80)
    p_cov2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cov2 = p_cov2.runs[0]
    r_cov2.font.name = "Georgia"
    r_cov2.font.size = Pt(22)
    r_cov2.font.bold = True

    p_cov2_meta = doc2.add_paragraph("Candidate ID: CAM-9921 • Michaelmas Term 2026")
    p_cov2_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov2_meta.runs[0].font.name = "Georgia"
    p_cov2_meta.runs[0].font.size = Pt(12)
    doc2.save(str(base_dir / "cover_with_custom_font_header.docx"))

    # -------------------------------------------------------------
    # 3. Official Cover (.pdf)
    # -------------------------------------------------------------
    pdf_doc = pymupdf.open()
    pdf_page = pdf_doc.new_page(width=595, height=842)  # A4
    pdf_page.draw_rect(pymupdf.Rect(40, 40, 555, 802), color=(0.1, 0.2, 0.5), width=2)
    pdf_page.insert_text(pymupdf.Point(120, 150), "IMPERIAL COLLEGE LONDON", fontsize=20, fontname="helv", color=(0.1, 0.2, 0.5))
    pdf_page.insert_text(pymupdf.Point(160, 180), "Department of Computing", fontsize=14, fontname="helv")
    pdf_page.insert_text(pymupdf.Point(100, 320), "AUTONOMOUS SYSTEM VERIFICATION REPORT", fontsize=16, fontname="helv", color=(0.15, 0.15, 0.15))
    pdf_page.insert_text(pymupdf.Point(100, 420), "Student: Sarah Jenkins (CID: 01849201)", fontsize=12, fontname="helv")
    pdf_page.insert_text(pymupdf.Point(100, 450), "Module: Safety-Critical Software (CO-528)", fontsize=12, fontname="helv")
    pdf_page.insert_text(pymupdf.Point(100, 480), "Date of Submission: October 2026", fontsize=12, fontname="helv")
    pdf_doc.save(str(base_dir / "official_cover.pdf"))
    pdf_doc.close()

    # -------------------------------------------------------------
    # 4. Sample AI Report with Markdown Tables & Artifacts (.docx)
    # -------------------------------------------------------------
    doc4 = docx.Document()
    doc4.add_paragraph("# 1. Introduction")
    doc4.add_paragraph("This is an **autonomous artificial intelligence** research paper examining high-dimensional feature spaces.")
    doc4.add_paragraph("## 1.1 Problem Motivation")
    doc4.add_paragraph("Deep learning models demonstrate unprecedented capabilities, but **rigorous formal verification** remains elusive.")
    doc4.add_paragraph("---")  # Markdown horizontal rule
    doc4.add_paragraph("Figure 1: Deep Neural Network Architecture Topology")
    doc4.add_paragraph("Below is the quantitative evaluation across test splits:")
    doc4.add_paragraph("| Model Architecture | Top-1 Accuracy | Inference Latency | Energy Consumption |")
    doc4.add_paragraph("|---|---|---|---|")
    doc4.add_paragraph("| Transformer Base | 84.2% | 18.2ms | 45W |")
    doc4.add_paragraph("| ConvNeXt Large | 86.7% | 22.4ms | 62W |")
    doc4.add_paragraph("| Hybrid Attention | 89.1% | 14.8ms | 38W |")
    doc4.add_paragraph("Table 1: Benchmark Evaluation Performance Matrix")
    doc4.add_paragraph("# 2. Conclusion")
    doc4.add_paragraph("The synthesized findings demonstrate significant empirical advantages in **throughput** and **accuracy**.")
    doc4.add_paragraph("# 3. References")
    doc4.add_paragraph("Smith, J. (2024). Foundations of Neural Verification. MIT Press.")
    doc4.save(str(base_dir / "sample_ai_report_with_tables.docx"))

    # -------------------------------------------------------------
    # 5. Plain Sample AI Report (.docx)
    # -------------------------------------------------------------
    doc5 = docx.Document()
    doc5.add_paragraph("# Introduction")
    doc5.add_paragraph("This assignment investigates modern distributed ledger architectures and Byzantine consensus protocols.")
    doc5.add_paragraph("# Methodology")
    doc5.add_paragraph("We systematically analyzed 40 peer-reviewed protocol specifications using rigorous comparative matrices.")
    doc5.add_paragraph("# Conclusion")
    doc5.add_paragraph("Proof-of-stake protocols demonstrate substantially lower energetic overhead with comparable finality guarantees.")
    doc5.add_paragraph("# References")
    doc5.add_paragraph("Nakamoto, S. (2008). Bitcoin: A Peer-to-Peer Electronic Cash System.")
    doc5.save(str(base_dir / "sample_ai_report_plain.docx"))

    # -------------------------------------------------------------
    # 6. Criteria APA (.pdf)
    # -------------------------------------------------------------
    crit_pdf = pymupdf.open()
    cp = crit_pdf.new_page(width=595, height=842)
    cp.insert_text(pymupdf.Point(50, 80), "FACULTY ASSESSMENT BRIEF & FORMATTING RULES", fontsize=14, fontname="helv", color=(0.1, 0.2, 0.5))
    text_apa = (
        "1. Paper Size: All assignments must be prepared on A4 paper.\n"
        "2. Typography: The manuscript body must be formatted in Times New Roman at 12 pt.\n"
        "3. Headings: Heading 1 must be 16 pt bold. Heading 2 must be 14 pt bold.\n"
        "4. Line Spacing: Body text must be double-spaced (2.0 spacing).\n"
        "5. Margins: Standard 1 inch margins on all sides (top, bottom, left, and right).\n"
        "6. Text Alignment: Body paragraphs must be left-aligned.\n"
        "7. Referencing Style: Adhere strictly to APA 7th edition referencing conventions."
    )
    y = 120
    for line in text_apa.split("\n"):
        cp.insert_text(pymupdf.Point(50, y), line, fontsize=10, fontname="helv")
        y += 24
    crit_pdf.save(str(base_dir / "criteria_apa.pdf"))
    crit_pdf.close()

    # -------------------------------------------------------------
    # 7. Criteria Harvard (.docx)
    # -------------------------------------------------------------
    crit_docx = docx.Document()
    crit_docx.add_paragraph("DEPARTMENTAL COURSEWORK SPECIFICATION")
    crit_docx.add_paragraph("Formatting Instructions:")
    crit_docx.add_paragraph("Font Family: Arial")
    crit_docx.add_paragraph("Body Font Size: 11 pt")
    crit_docx.add_paragraph("Line Spacing: 1.5 line spacing")
    crit_docx.add_paragraph("Main Headings: 14 pt bold")
    crit_docx.add_paragraph("Margins: Standard 1 inch on all sides")
    crit_docx.add_paragraph("Alignment: Justified")
    crit_docx.add_paragraph("Referencing Style: Harvard citation format")
    crit_docx.save(str(base_dir / "criteria_harvard.docx"))

    # -------------------------------------------------------------
    # 8. Marking Scheme (.pdf)
    # -------------------------------------------------------------
    rubric_pdf = pymupdf.open()
    rp = rubric_pdf.new_page(width=595, height=842)
    rp.insert_text(pymupdf.Point(50, 80), "MODULE MARKING SCHEME & ASSESSMENT RUBRIC", fontsize=14, fontname="helv", color=(0.2, 0.1, 0.5))
    r_text = (
        "Component 1: Theoretical Understanding & Analysis — 40% (40 Marks)\n"
        "Component 2: Research Methodology & Evidence Base — 30% (30 Marks)\n"
        "Component 3: Academic Presentation & Structure — 15% (15 Marks)\n"
        "Component 4: Referencing Accuracy & Integrity — 15% (15 Marks)\n"
        "Total Maximum Available: 100 Marks (100%)"
    )
    y = 120
    for line in r_text.split("\n"):
        rp.insert_text(pymupdf.Point(50, y), line, fontsize=11, fontname="helv")
        y += 28
    rubric_pdf.save(str(base_dir / "marking_scheme.pdf"))
    rubric_pdf.close()

    logo_path.unlink(missing_ok=True)
    print(f"All 8 test fixtures created in: {base_dir}")


if __name__ == "__main__":
    import sys
    target = Path("apps/api/tests/fixtures")
    create_fixtures(target)
    # Also mirror in tests/fixtures for root test execution
    create_fixtures(Path("tests/fixtures"))
